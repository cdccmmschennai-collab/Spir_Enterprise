"""
Saipem SPIR output workbook.

Header (rows 1-3) mirrors the form's hierarchy with merged header cells:

    row 1  band   QUANTITY FOR 2 YEARS OPERATION | CAPITAL SPARES | …
    row 2  field  RECOMMEND. By VENDOR | ENGINEER. RECOMMEND. | PRICE | …
    row 3  sub    Total quantity | Quantity for PM | Quantity for CM | UNIT | …

A column without a band or sub-header has its label merged down over the
empty header rows. Data starts on row 4; data cells are never merged, so
every record keeps its own value. Header labels keep the source
capitalization; data strings are already uppercased by the extractor.

Generic-only behaviour (S.NO, ERROR, SPIR metadata/char-limit rows) is not
applied — the generic services/excel_builder.build_xlsx is left untouched.
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

from spir_dynamic.extraction.saipem.profile import SaipemColumn, header_parts
from spir_dynamic.services.excel_builder import _register_styles

_HEADER_ROWS = 3
_HDR_HEIGHT = 30


def build_saipem_xlsx(
    rows: list[dict[str, Any]],
    columns: list[SaipemColumn],
    band_labels: dict[str, str],
    spir_no: str = "",
) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    safe_title = re.sub(r'[\\/*?:\[\]]+', ' ', spir_no or "SPIR Extraction").strip()
    ws.title = safe_title[:31] or "SPIR Extraction"
    _register_styles(wb)

    parts = [header_parts(c, band_labels) for c in columns]
    _write_header(ws, columns, parts)

    for r_idx, row in enumerate(rows, start=_HEADER_ROWS + 1):
        for c_idx, col in enumerate(columns, start=1):
            v = row.get(col.id)
            if v is None:
                continue
            cell = ws.cell(r_idx, c_idx, value=v)
            if col.identifier and isinstance(v, str):
                cell.number_format = "@"
            elif isinstance(v, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss" if v.time() != datetime.min.time() else "yyyy-mm-dd"
            elif isinstance(v, date):
                cell.number_format = "yyyy-mm-dd"

    for c_idx, col in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(c_idx)].width = col.width
    for r in range(1, _HEADER_ROWS + 1):
        ws.row_dimensions[r].height = _HDR_HEIGHT
    ws.freeze_panes = ws.cell(_HEADER_ROWS + 1, 1)
    if columns:
        ws.auto_filter.ref = f"A{_HEADER_ROWS}:{get_column_letter(len(columns))}{_HEADER_ROWS}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_header(ws, columns: list[SaipemColumn], parts) -> None:
    n = len(columns)
    for c in range(1, n + 1):
        for r in range(1, _HEADER_ROWS + 1):
            ws.cell(r, c).style = "spir_hdr"

    # Row 1: contiguous columns sharing a band
    for start, end in _runs([columns[i].band if parts[i][0] else None for i in range(n)]):
        band = parts[start][0]
        if band:
            _merge_write(ws, 1, start + 1, 1, end + 1, band)

    # Rows 2-3 (or 1-3 without a band): field label, then sub-headers
    field_keys = [(columns[i].band, parts[i][1]) for i in range(n)]
    for start, end in _runs(field_keys):
        top = 2 if parts[start][0] else 1
        has_sub = any(parts[i][2] for i in range(start, end + 1))
        field_bottom = _HEADER_ROWS - 1 if has_sub else _HEADER_ROWS
        _merge_write(ws, top, start + 1, field_bottom, end + 1, parts[start][1])
        if has_sub:
            for i in range(start, end + 1):
                ws.cell(_HEADER_ROWS, i + 1, value=parts[i][2])


def _runs(keys: list) -> list[tuple[int, int]]:
    """(start, end) index pairs of maximal runs of equal keys."""
    runs: list[tuple[int, int]] = []
    i = 0
    while i < len(keys):
        j = i
        while j + 1 < len(keys) and keys[j + 1] == keys[i] and keys[i] is not None:
            j += 1
        runs.append((i, j))
        i = j + 1
    return runs


def _merge_write(ws, r1: int, c1: int, r2: int, c2: int, value: Any) -> None:
    ws.cell(r1, c1, value=value)
    if (r1, c1) != (r2, c2):
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
