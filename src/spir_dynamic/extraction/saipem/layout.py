"""
Saipem SPIR form layout — every position is found from source labels and
merged ranges, never from fixed rows or column letters.

Structure of the form (as found in EX2515011 and its Example sheet):

    label column  ITEM TAG No. / Free Field / SERIAL No. / MODEL OR TYPE /
                  EQUIPMENT SUB ITEM No. / No. OF UNITS / QUANTITY OF PARTS
                  INSTALLED PER ONE UNIT … (the last one is merged over exactly
                  the spare-row span)
    code column   A, B, C, D, E … (skipped)
    equipment     one column per equipment group, between the code column and
    columns       the TOTAL QUANTITY OF IDENTICAL PARTS INSTALLED column
    matrix        header band row (SPARE PARTS IDENTIFICATION, QUANTITY FOR 2
                  YEARS OPERATION, CAPITAL SPARES), then 1-3 header rows whose
                  merged ranges form a tree: field → sub-header → qualifier

The header tree is matched to SAIPEM_COLUMNS ids by label. A matrix column
that matches no known field is kept as an extra column with its full source
header path, so nothing in the matrix is dropped silently.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Callable, Optional

from openpyxl.utils import get_column_letter

from spir_dynamic.extraction.saipem.profile import (
    BAND_CAPITAL,
    BAND_EQUIPMENT,
    BAND_OPERATING,
    BAND_SPARE,
    COLUMN_BY_ID,
)
from spir_dynamic.extraction.saipem.values import MergeIndex, display_label, norm_label


@dataclass(frozen=True)
class Label:
    row: int
    col: int
    max_row: int
    max_col: int
    text: str        # display form
    norm: str        # comparable form

    def covers_col(self, col: int) -> bool:
        return self.col <= col <= self.max_col


# Equipment label predicates (normalized text)
def _starts(prefix: str) -> Callable[[str], bool]:
    return lambda n: n.startswith(prefix)


EQUIPMENT_LABELS: dict[str, Callable[[str], bool]] = {
    "tag": _starts("ITEM TAG NO"),
    "free_field": _starts("FREE FIELD"),
    "serial": _starts("SERIAL NO"),
    "model": _starts("MODEL OR TYPE"),
    "sub_item": _starts("EQUIPMENT SUB ITEM NO"),
    "units": _starts("NO. OF UNITS"),
    "per_unit": _starts("QUANTITY OF PARTS INSTALLED PER ONE UNIT"),
}
IS_TOTAL_INSTALLED = _starts("TOTAL QUANTITY OF IDENTICAL PARTS INSTALLED")
IS_SPIR_POS = _starts("SPIR POS")
IS_DESCRIPTION = _starts("DESCRIPTION OF PARTS")
IS_SPARE_BAND = _starts("SPARE PARTS IDENTIFICATION")

# Header-block labels whose right-hand neighbour is another label, not a value
_KNOWN_HEADER_LABELS = frozenset({
    "SPIR NO.", "REV", "SHEET NO.", "OF", "PROJECT", "OWNER", "ENGINEERED BY",
    "CONTRACT NO.", "MAIN P.O. DESCRIPTION", "MAIN PURCHASE ORDER NO.",
    "VENDOR'S NAME", "ADDRESS", "PERSON IN CHARGE", "SECTION",
    "PHONE NO. / FAX NO.", "E-MAIL ADDRESS", "VENDOR'S ISSUE DATE",
    "CHECKED BY", "CHECKED BY :", "APPROVED BY", "REVISION NO.", "DESCRIPTION",
    "SIGNATURE", "DATE",
})

_FOOTER_MARKERS = ("SUB-TOTAL", "TOTAL :", "TOTAL:")
_CURRENCY_CODE_RE = re.compile(r"\(\s*([A-Z]{3})\s*\)")
_PM_RE = re.compile(r"\bPM\b")
_CM_RE = re.compile(r"\bCM\b")


@dataclass
class MatrixColumn:
    col: int
    band_key: Optional[str]
    band_label: Optional[str]
    path: list[str]          # display labels, top-down
    field_id: Optional[str] = None


@dataclass
class SaipemLayout:
    sheet: str
    merges: MergeIndex
    label_col: int
    equipment_rows: dict[str, tuple[int, int]]   # field → (first row, last row) of its label
    equipment_cols: list[int]
    total_col: Optional[int]
    pos_col: int
    data_first: int
    data_last: int
    columns: dict[str, int] = field(default_factory=dict)          # field id → column
    extra_columns: list[MatrixColumn] = field(default_factory=list)
    band_labels: dict[str, str] = field(default_factory=dict)       # band key → source label
    currency: Optional[str] = None
    labels: list[Label] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def find_label(self, pred: Callable[[str], bool]) -> Optional[Label]:
        return next((lab for lab in self.labels if pred(lab.norm)), None)


def scan_labels(ws, merges: MergeIndex, max_row: Optional[int] = None,
                max_col: Optional[int] = None) -> list[Label]:
    """Every text cell (merged-range anchors only), in reading order.

    Bounded by the sheet's own dimensions so no cells are created.
    """
    max_row = min(max_row or ws.max_row or 0, ws.max_row or 0)
    max_col = min(max_col or ws.max_column or 0, ws.max_column or 0)
    labels: list[Label] = []
    if not max_row or not max_col:
        return labels
    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
        for cell in row:
            v = cell.value
            if not isinstance(v, str) or not v.strip():
                continue
            r1, c1, r2, c2 = merges.bounds(cell.row, cell.column)
            labels.append(Label(r1, c1, r2, c2, display_label(v), norm_label(v)))
    return labels


def band_key_for(norm: str) -> Optional[str]:
    if IS_SPARE_BAND(norm):
        return BAND_SPARE
    if "CAPITAL" in norm:
        return BAND_CAPITAL
    if "OPERATION" in norm or "OPERATING" in norm:
        return BAND_OPERATING
    if norm.startswith("EQUIPMENT IDENTIFICATION"):
        return BAND_EQUIPMENT
    return None


def resolve_layout(ws) -> Optional[SaipemLayout]:
    """Locate every Saipem structure on the sheet; None if the form is incomplete."""
    merges = MergeIndex(ws)
    labels = scan_labels(ws, merges)

    def first(pred, where=lambda lab: True) -> Optional[Label]:
        return next((lab for lab in labels if pred(lab.norm) and where(lab)), None)

    tag = first(EQUIPMENT_LABELS["tag"])
    total = first(IS_TOTAL_INSTALLED)
    pos = first(IS_SPIR_POS)
    if tag is None or pos is None:
        return None
    label_col = tag.col
    total_col = total.col if total is not None and total.col > label_col else None

    equipment_rows: dict[str, tuple[int, int]] = {}
    for key, pred in EQUIPMENT_LABELS.items():
        lab = first(pred, lambda lab: lab.covers_col(label_col) and lab.col < pos.col)
        if lab is not None:
            equipment_rows[key] = (lab.row, lab.max_row)

    layout = SaipemLayout(
        sheet=ws.title, merges=merges, label_col=label_col,
        equipment_rows=equipment_rows, equipment_cols=[],
        total_col=total_col, pos_col=pos.col, data_first=0, data_last=-1,
        labels=labels,
    )

    header_top, header_bottom = pos.row, pos.max_row
    layout.equipment_cols = _equipment_columns(ws, layout, right_limit=total_col or pos.col)
    layout.data_first, layout.data_last = _data_span(ws, layout, header_bottom)
    _map_matrix(layout, header_top, header_bottom)
    return layout


def _equipment_columns(ws, layout: SaipemLayout, right_limit: int) -> list[int]:
    """Columns between the label column and the total column that hold equipment.

    The code column (single letters A, B, C … beside each label) is skipped.
    """
    rows = layout.equipment_rows
    tag_rows = rows.get("tag")
    code_probe = [rows[k][0] for k in ("tag", "serial", "model") if k in rows]
    cols: list[int] = []
    for c in range(layout.label_col + 1, right_limit):
        probes = [norm_label(ws.cell(r, c).value) for r in code_probe]
        if probes and all(len(p) == 1 and p.isalpha() for p in probes):
            continue  # code column
        if tag_rows and any(
            layout.merges.anchor(r, c) == (r, c) and norm_label(ws.cell(r, c).value)
            for r in range(tag_rows[0], tag_rows[1] + 1)
        ):
            cols.append(c)
    return cols


def _data_span(ws, layout: SaipemLayout, header_bottom: int) -> tuple[int, int]:
    """Spare rows: the span of the merged 'QUANTITY OF PARTS INSTALLED PER ONE
    UNIT' label; without that merge, from below the headers to the footer."""
    per_unit = layout.equipment_rows.get("per_unit")
    if per_unit and per_unit[1] > per_unit[0]:
        return per_unit
    start = max(header_bottom, layout.equipment_rows.get("units", (0, 0))[1]) + 1
    for lab in layout.labels:
        if lab.row >= start and lab.norm in _FOOTER_MARKERS:
            return start, lab.row - 1
    return start, ws.max_row or start


def _map_matrix(layout: SaipemLayout, header_top: int, header_bottom: int) -> None:
    labels = layout.labels
    band_row_label = next(
        (lab for lab in labels if IS_SPARE_BAND(lab.norm) and lab.row < header_top), None
    )
    band_row = band_row_label.row if band_row_label else header_top - 1
    bands = [lab for lab in labels if lab.row == band_row and lab.max_col >= layout.pos_col]
    for lab in bands:
        key = band_key_for(lab.norm)
        if key and key not in layout.band_labels:
            layout.band_labels[key] = lab.text
    eq_band = next((lab for lab in labels if lab.norm.startswith("EQUIPMENT IDENTIFICATION")), None)
    if eq_band:
        layout.band_labels.setdefault(BAND_EQUIPMENT, eq_band.text)

    headers = [
        lab for lab in labels
        if header_top <= lab.row <= header_bottom and lab.col >= layout.pos_col
    ]
    data_cols = sorted({lab.col for lab in headers})
    for c in data_cols:
        path_labels = sorted((lab for lab in headers if lab.covers_col(c)), key=lambda lab: lab.row)
        band = next((lab for lab in bands if lab.covers_col(c)), None)
        mc = MatrixColumn(
            col=c,
            band_key=band_key_for(band.norm) if band else None,
            band_label=band.text if band else None,
            path=[lab.text for lab in path_labels],
        )
        norm_path = [lab.norm for lab in path_labels]
        _note_currency(layout, norm_path)
        fid = _match_field(mc.band_key, norm_path)
        if fid is not None and fid in COLUMN_BY_ID and fid not in layout.columns:
            mc.field_id = fid
            layout.columns[fid] = c
        else:
            layout.extra_columns.append(mc)
            layout.warnings.append(
                f"Sheet '{layout.sheet}' column {get_column_letter(c)}: header "
                f"'{' / '.join(filter(None, [mc.band_label, *mc.path]))}' is not a standard "
                "Saipem field; exported as an additional column."
            )


def _note_currency(layout: SaipemLayout, norm_path: list[str]) -> None:
    if layout.currency is not None or not norm_path or not norm_path[0].startswith("PRICE"):
        return
    for n in norm_path:
        if n.startswith("CURRENCY"):
            m = _CURRENCY_CODE_RE.search(n)
            if m:
                layout.currency = m.group(1)


def _match_field(band: Optional[str], p: list[str]) -> Optional[str]:
    """Field id for a matrix column's header path, or None when unknown."""
    if not p:
        return None
    first, rest = p[0], p[1:]
    if IS_SPIR_POS(first):
        return "spir_pos_no"
    if IS_DESCRIPTION(first):
        return "description_of_parts"
    if first.startswith("MATERIAL"):
        return "material"
    if first.startswith("ORIGINAL MAKERS NAME"):
        return "original_makers_name"
    if first.startswith("ORIGINAL MAKERS PART"):
        return "original_makers_part_no"
    if first.startswith("RELEVANT DOCUMENT"):
        last = rest[-1] if rest else ""
        if last.startswith("DRAWING"):
            return "rel_doc_drawing_no"
        if last == "PAGE":
            return "rel_doc_page"
        if last.startswith("POS"):
            return "rel_doc_pos_no"
        return None
    if first.startswith("VENDOR CODE"):
        return "vendor_code_no"
    if "MIAP" in first or first.startswith("BUSINESS UNIT CODIFICATION"):
        return "miap_code"
    if first.startswith("UNIT OF MEAS"):
        return "unit_of_meas"
    if first.startswith("RECOMMEND") and "VENDOR" in first:
        if band == BAND_OPERATING:
            if not rest:
                return "op2y_vendor_total"
            if _PM_RE.search(rest[-1]):
                return "op2y_vendor_pm"
            if _CM_RE.search(rest[-1]):
                return "op2y_vendor_cm"
            return None
        if band == BAND_CAPITAL and not rest:
            return "cap_vendor"
        return None
    for prefix, suffix in (("ENGINEER", "engineer"), ("OPERATION RECOMMEND", "operation"),
                           ("PURCHASED QUANTITY", "purchased")):
        if first.startswith(prefix):
            if band == BAND_OPERATING:
                return f"op2y_{suffix}"
            if band == BAND_CAPITAL:
                return f"cap_{suffix}"
            return None
    if first.startswith("WEIGHT"):
        return "weight"
    if first.startswith("CHARACTERISTIC DIMENSION"):
        return "characteristic_dimension"
    if first.startswith("OBSOLESCENCE"):
        return "obsolescence"
    if first == "G/S":
        return "gs"
    if first.startswith("SHELF LIFE"):
        return "shelf_life"
    if first.startswith("PRICE"):
        if "UNIT" in rest:
            return "price_unit"
        if "AMOUNT" in rest:
            return "price_amount"
        if rest and rest[-1].startswith("CURRENCY"):
            return "price_currency"
        return None
    if first.startswith("DELIVERY TIME"):
        return "delivery_time_weeks"
    return None


def labelled_value_cell(ws, layout: SaipemLayout, pred: Callable[[str], bool]):
    """The value cell to the right of a header-block label (e.g. SPIR No., VENDOR'S NAME).

    Skips empty cells; stops at another known label so an empty value never
    borrows its neighbour's label. Returns None when no value is present.
    """
    lab = layout.find_label(pred)
    if lab is None:
        return None
    c = lab.max_col + 1
    limit = min(ws.max_column or 0, lab.max_col + 16)
    while c <= limit:
        r1, c1, _, c2 = layout.merges.bounds(lab.row, c)
        if r1 == lab.row:
            cell = ws.cell(r1, c1)
            v = cell.value
            if v is not None and not (isinstance(v, str) and not v.strip()):
                if isinstance(v, str) and norm_label(v) in _KNOWN_HEADER_LABELS:
                    return None
                return cell
        c = c2 + 1
    return None
