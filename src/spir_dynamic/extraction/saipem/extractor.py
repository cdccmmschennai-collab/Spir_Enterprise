"""
Saipem SPIR extraction — assembles output rows from the detected form sheets.

Row convention (same as the generic COLUMN_HEADERS path, strategies/columnar.py):
for each equipment group, for each expanded tag that has applicable spares —
  1. one equipment row: SPIR No., tag, equipment metadata, No. OF UNITS and,
     as DESCRIPTION OF PARTS, the form's labelled DESCRIPTION (the equipment's)
  2. one row per applicable SPIR line: equipment metadata repeated, the tag's
     installed quantity per unit, and the line's spare data. No. OF UNITS is
     only on the equipment row, as eqpt_qty is in the generic output.
Spare quantities are repeated unchanged on every applicable tag's row.

SPIR No. and revision come from the sheet's own SPIR No. / REV fields (the
filename is never used); they are joined as '<SPIR No.>-REV-<rev>', the same
convention the generic pipeline uses.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from typing import Any, Optional

from openpyxl.utils import get_column_letter

from spir_dynamic.extraction.saipem.equipment import EquipmentGroup, read_equipment_groups
from spir_dynamic.extraction.saipem.layout import (
    MatrixColumn,
    SaipemLayout,
    labelled_value_cell,
    resolve_layout,
)
from spir_dynamic.extraction.saipem.matrix import SpareLine, read_spare_lines
from spir_dynamic.extraction.saipem.profile import (
    COLUMN_BY_ID,
    COMPUTED_FIELD_IDS,
    DEFAULT_BAND_LABELS,
    MATRIX_FIELD_IDS,
    SAIPEM_COLUMNS,
    SaipemColumn,
)
from spir_dynamic.extraction.saipem.values import clean_value, identifier_text

_VENDOR_LABELS = {
    "vendor_name": lambda n: n == "VENDOR'S NAME",
    "vendor_address": lambda n: n == "ADDRESS",
    "vendor_phone_fax": lambda n: n.startswith("PHONE NO"),
    "vendor_email": lambda n: n.startswith("E-MAIL"),
}


@dataclass
class SaipemExtraction:
    rows: list[dict[str, Any]]
    columns: list[SaipemColumn]
    band_labels: dict[str, str]
    spir_no: str
    vendor: dict[str, Any] = field(default_factory=dict)
    sheets: list[str] = field(default_factory=list)
    total_tags: int = 0
    spare_items: int = 0
    warnings: list[str] = field(default_factory=list)


def extract_saipem(wb, sheet_names: list[str]) -> SaipemExtraction:
    columns: list[SaipemColumn] = list(SAIPEM_COLUMNS)
    band_labels = dict(DEFAULT_BAND_LABELS)
    source_bands: dict[str, str] = {}
    rows: list[dict[str, Any]] = []
    warnings: list[str] = []
    spir_no = ""
    vendor: dict[str, Any] = {}
    tags_seen: set[str] = set()
    lines_used: set[tuple[str, int]] = set()
    extra_ids: dict[tuple[str, int], str] = {}

    for name in sheet_names:
        ws = wb[name]
        layout = resolve_layout(ws)
        if layout is None:
            warnings.append(f"Sheet '{name}': Saipem form markers incomplete; sheet skipped.")
            continue
        warnings.extend(layout.warnings)
        # The workbook's own group headings replace the defaults (first sheet wins).
        for key, label in layout.band_labels.items():
            if key not in source_bands:
                source_bands[key] = label
                band_labels[key] = label

        sheet_spir = _spir_no(ws, layout)
        if sheet_spir and not spir_no:
            spir_no = sheet_spir
        elif sheet_spir and sheet_spir != spir_no:
            warnings.append(f"Sheet '{name}': SPIR No. '{sheet_spir}' differs from '{spir_no}'.")
        # The form's DESCRIPTION field (revision block, e.g. "CENTRIFUGAL PUMPS")
        # describes the equipment; it fills the description on each tag's first row.
        equipment_description = _labelled(ws, layout, lambda n: n == "DESCRIPTION", "description_of_parts")
        sheet_vendor = {fid: _labelled(ws, layout, pred, fid) for fid, pred in _VENDOR_LABELS.items()}
        for fid, v in sheet_vendor.items():
            if v is not None:
                vendor.setdefault(fid, v)

        _missing_fields(layout, warnings)
        for mc in layout.extra_columns:
            col_def = _extra_column(mc, name, band_labels)
            extra_ids[(name, mc.col)] = col_def.id
            columns.append(col_def)

        groups = read_equipment_groups(ws, layout)
        lines = read_spare_lines(ws, layout)
        for g in groups:
            warnings.extend(g.warnings)
        _warn_unassociated(name, lines, groups, warnings)

        for g in groups:
            applicable = [ln for ln in lines if ln.per_unit.get(g.col) is not None]
            for i, tag in enumerate(g.tags):
                if not applicable:
                    warnings.append(f"Tag '{tag}' has no applicable spare rows; not exported.")
                    continue
                tags_seen.add(tag.upper())
                base = {
                    "spir_no": None,   # filled once the SPIR No. is final
                    "item_tag_no": tag,
                    "serial_no": g.serials[i] if i < len(g.serials) else None,
                    **g.meta,
                    **sheet_vendor,
                }
                # Equipment row: like the Standard header row, its description
                # is the equipment's (the form's labelled DESCRIPTION field).
                rows.append({**base, "no_of_units": g.no_of_units,
                             "description_of_parts": equipment_description})
                for ln in applicable:
                    lines_used.add((name, ln.row))
                    row = {**base,
                           **{fid: ln.values.get(fid) for fid in MATRIX_FIELD_IDS | COMPUTED_FIELD_IDS},
                           "qty_installed_per_unit": ln.per_unit.get(g.col),
                           "total_qty_identical_installed": ln.total_installed}
                    for col, v in ln.extras.items():
                        row[extra_ids[(name, col)]] = v
                    rows.append(row)

    for row in rows:
        row["spir_no"] = spir_no or None
    _uppercase(rows)
    return SaipemExtraction(
        rows=rows, columns=columns, band_labels=band_labels,
        spir_no=spir_no.upper(), vendor=vendor, sheets=list(sheet_names),
        total_tags=len(tags_seen), spare_items=len(lines_used),
        warnings=list(dict.fromkeys(warnings)),
    )


def _spir_no(ws, layout: SaipemLayout) -> str:
    """'<SPIR No.>-REV-<rev>' from the sheet's labelled cells (REV, else REVISION No.)."""
    cell = labelled_value_cell(ws, layout, lambda n: n == "SPIR NO.")
    number = identifier_text(cell) if cell is not None else None
    if not number:
        return ""
    rev_cell = (labelled_value_cell(ws, layout, lambda n: n == "REV")
                or labelled_value_cell(ws, layout, lambda n: n == "REVISION NO."))
    rev = identifier_text(rev_cell) if rev_cell is not None else None
    return f"{number}-REV-{rev}" if rev else number


def _labelled(ws, layout: SaipemLayout, pred, fid: str) -> Any:
    cell = labelled_value_cell(ws, layout, pred)
    if cell is None:
        return None
    return identifier_text(cell) if COLUMN_BY_ID[fid].identifier else clean_value(cell.value)


def _missing_fields(layout: SaipemLayout, warnings: list[str]) -> None:
    missing = [COLUMN_BY_ID[f] for f in sorted(MATRIX_FIELD_IDS) if f not in layout.columns]
    if "price_currency" not in layout.columns and layout.currency is None:
        missing.append(COLUMN_BY_ID["price_currency"])
    for col in missing:
        label = " / ".join(filter(None, [col.field, col.sub]))
        warnings.append(f"Sheet '{layout.sheet}': field '{label}' not found in the matrix headers; left blank.")
    for key in ("tag", "serial", "model", "units", "per_unit"):
        if key not in layout.equipment_rows:
            warnings.append(f"Sheet '{layout.sheet}': equipment label '{key}' not found; left blank.")


def _extra_column(mc: MatrixColumn, sheet: str, band_labels: dict[str, str]) -> SaipemColumn:
    band_key = None
    if mc.band_label:
        band_key = mc.band_key or f"extra:{mc.band_label}"
        band_labels.setdefault(band_key, mc.band_label)
    field_label = mc.path[0] if mc.path else get_column_letter(mc.col)
    sub = " / ".join(mc.path[1:]) or None
    return SaipemColumn(
        id=f"extra__{sheet}__{get_column_letter(mc.col)}",
        field=field_label, sub=sub, band=band_key,
    )


def _warn_unassociated(sheet: str, lines: list[SpareLine], groups: list[EquipmentGroup],
                       warnings: list[str]) -> None:
    group_cols = {g.col for g in groups if g.tags}
    for ln in lines:
        if not any(ln.per_unit.get(c) is not None for c in group_cols):
            pos = ln.values.get("spir_pos_no") or f"row {ln.row}"
            warnings.append(
                f"Sheet '{sheet}': SPIR POS. No. {pos} has no installed quantity for any "
                "tagged equipment column; not associated with a tag."
            )


def _uppercase(rows: list[dict[str, Any]]) -> None:
    for row in rows:
        for k, v in row.items():
            if isinstance(v, str):
                row[k] = v.upper()
