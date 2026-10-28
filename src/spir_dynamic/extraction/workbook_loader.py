"""
Format-aware workbook loading — the single place a source file becomes a workbook.

.xlsx / .xlsm keep the exact openpyxl call the pipeline has always used.

.xls (BIFF8 in an OLE2 compound file) is read with xlrd and copied into an
in-memory openpyxl Workbook, so every downstream analyzer and strategy sees the
same interface (ws.cell().value, max_row/max_column, merged_cells, sheet_state,
title) it sees for .xlsx. Values are converted the way openpyxl's own .xlsx
reader converts them (see _cell_value). Legacy Forms checkboxes, which xlrd
does not expose, are parsed from the BIFF stream and attached as
wb._spir_form_checkboxes = {sheet title: (FormCheckbox, ...)} for SPIR-type
detection (header_detector._detect_spir_type, Phase 0a).

Every workbook carries _spir_raw_bytes / _spir_raw_path (the VML checkbox
reader re-opens the .xlsx ZIP from them); for .xls both are None.
"""
from __future__ import annotations

import io
import logging
import struct
from pathlib import Path
from typing import Union

import openpyxl
import xlrd
from openpyxl.styles.numbers import is_date_format, is_timedelta_format
from openpyxl.utils.datetime import MAC_EPOCH, WINDOWS_EPOCH, from_excel
from xlrd.compdoc import CompDoc, CompDocError

from spir_dynamic.extraction.file_validator import (
    LEGACY_XLS_EXTENSION,
    OOXML_EXTENSIONS,
    ValidationError,
    file_extension,
    unsupported_extension_message,
)
from spir_dynamic.extraction.xls_form_controls import read_xls_checkboxes

log = logging.getLogger(__name__)

# openpyxl reads an integral number stored as "13" as int and "1E+20" as float.
# Excel writes integral doubles below 1e15 without exponent or decimal point.
_INT_LIMIT = 1e15

_VISIBILITY = {0: "visible", 1: "hidden", 2: "veryHidden"}

# What xlrd raises for a damaged compound file / BIFF stream (deterministic for
# a given file, so reported as a validation error, never retried).
_XLS_PARSE_ERRORS = (
    CompDocError, struct.error, AssertionError, IndexError, KeyError,
    ValueError, UnicodeDecodeError,
)


def load_workbook_any(source: Union[Path, bytes], filename: str) -> openpyxl.Workbook:
    """
    Open `source` (a temp-file Path or raw bytes) according to `filename`'s extension.

    Raises ValidationError for unsupported extensions and for .xls files xlrd
    cannot read (encrypted, malformed). Callers validate signatures first
    (file_validator.validate_file).
    """
    ext = file_extension(filename)
    if ext in OOXML_EXTENSIONS:
        return _load_ooxml(source)
    if ext == LEGACY_XLS_EXTENSION:
        raw = source.read_bytes() if isinstance(source, Path) else source
        return load_xls_bytes(raw)
    raise ValidationError(unsupported_extension_message(ext))


def _load_ooxml(source: Union[Path, bytes]) -> openpyxl.Workbook:
    """The pipeline's original openpyxl load, unchanged."""
    # Not read_only — merged_cells.ranges used in column_mapper.
    # keep_links=False skips external link parsing.
    if isinstance(source, Path):
        wb = openpyxl.load_workbook(str(source), data_only=True, keep_links=False)
        wb._spir_raw_bytes = None
        # VML form-control detection re-opens the ZIP from the temp-file path.
        wb._spir_raw_path = str(source)
    else:
        wb = openpyxl.load_workbook(io.BytesIO(source), data_only=True, keep_links=False)
        wb._spir_raw_bytes = source
        wb._spir_raw_path = None
    return wb


def load_xls_bytes(raw: bytes) -> openpyxl.Workbook:
    """Read a BIFF .xls workbook into an equivalent in-memory openpyxl Workbook."""
    try:
        # formatting_info=True: merged ranges, blank (formatted) cells and the
        # number formats needed to tell dates from numbers.
        book = xlrd.open_workbook(file_contents=raw, formatting_info=True, logfile=_NullLog())
    except xlrd.XLRDError as exc:
        if "encrypted" in str(exc).lower():
            raise ValidationError(
                "This .xls workbook is password-protected. Remove the password in "
                "Excel and upload it again."
            ) from exc
        raise ValidationError(f"Cannot read this .xls workbook: {exc}") from exc
    except _XLS_PARSE_ERRORS as exc:
        raise ValidationError(f"Cannot read this .xls workbook — the file is damaged: {exc}") from exc

    try:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        epoch = MAC_EPOCH if book.datemode == 1 else WINDOWS_EPOCH
        for sheet in book.sheets():
            ws = wb.create_sheet(title=sheet.name)
            ws.sheet_state = _VISIBILITY.get(sheet.visibility, "visible")
            for r in range(sheet.nrows):
                for c, cell in enumerate(sheet.row(r)):
                    if cell.ctype == xlrd.XL_CELL_EMPTY:
                        continue
                    # BLANK (formatted, no value) cells are created too, as
                    # openpyxl does for styled empty cells — same max_row/max_column.
                    ws.cell(row=r + 1, column=c + 1, value=_cell_value(book, cell, epoch))
            # After the values: merging turns non-anchor cells into MergedCell,
            # exactly as openpyxl's .xlsx reader does.
            for rlo, rhi, clo, chi in sheet.merged_cells:
                ws.merge_cells(start_row=rlo + 1, end_row=rhi, start_column=clo + 1, end_column=chi)

        wb._spir_raw_bytes = None
        wb._spir_raw_path = None
        wb._spir_form_checkboxes = _form_checkboxes(raw, wb)
        return wb
    finally:
        book.release_resources()


def _cell_value(book, cell, epoch):
    """One xlrd cell → the value openpyxl(data_only=True) would give for .xlsx."""
    ctype, value = cell.ctype, cell.value
    if ctype == xlrd.XL_CELL_TEXT:
        return value
    if ctype in (xlrd.XL_CELL_NUMBER, xlrd.XL_CELL_DATE):
        fmt = _number_format(book, cell.xf_index)
        is_date = is_date_format(fmt) if fmt is not None else ctype == xlrd.XL_CELL_DATE
        if is_date:
            try:
                return from_excel(value, epoch, timedelta=bool(fmt) and is_timedelta_format(fmt))
            except (OverflowError, ValueError):
                return "#VALUE!"  # openpyxl's treatment of an out-of-range date serial
        if float(value).is_integer() and abs(value) < _INT_LIMIT:
            return int(value)
        return value
    if ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(value)
    if ctype == xlrd.XL_CELL_ERROR:
        return xlrd.error_text_from_code.get(value, "#VALUE!")
    return None  # XL_CELL_BLANK


def _number_format(book, xf_index):
    """The number-format string of an XF, or None when unavailable."""
    try:
        xf = book.xf_list[xf_index]
        return book.format_map[xf.format_key].format_str
    except (IndexError, KeyError, TypeError, AttributeError):
        return None


def _form_checkboxes(raw: bytes, wb: openpyxl.Workbook) -> dict | None:
    """{sheet title: (FormCheckbox, ...)} or None when the controls cannot be read."""
    try:
        doc = CompDoc(raw, logfile=_NullLog())
        stream = doc.get_named_stream("Workbook")
    except Exception as exc:
        log.warning("xls form controls: cannot open Workbook stream: %s", exc)
        return None
    if not stream:
        return None
    per_sheet = read_xls_checkboxes(stream)
    if per_sheet is None:
        return None
    if len(per_sheet) != len(wb.worksheets):
        log.warning(
            "xls form controls: %d worksheet substreams vs %d sheets — ignoring controls",
            len(per_sheet), len(wb.worksheets),
        )
        return None
    return {ws.title: tuple(boxes) for ws, boxes in zip(wb.worksheets, per_sheet)}


class _NullLog:
    """xlrd writes warnings to a file-like logfile (stdout by default)."""

    def write(self, *_args) -> None:
        pass

    def flush(self) -> None:
        pass
