"""
"Refer column 12A" in EQPT MAKE/MODEL/SR NO resolves to the column numbered
12A in the sheet's column-number row; equipment rows take the named value
after the slash, else the tag's most common resolved value.
"""
import openpyxl
import pytest

from spir_dynamic.extraction.column_reference import (
    find_column_codes,
    parse_column_reference,
    resolve_column_references,
)
from spir_dynamic.models.sheet_profile import SheetProfile


@pytest.mark.parametrize("value, expected", [
    ("Refer column 12A", ("12A", None)),
    ("Refer column 12A/YOKOGAWA", ("12A", "YOKOGAWA")),
    ("REFER TO COL. 9", ("9", None)),
    ("refer col no 10a / Acme Ltd", ("10A", "Acme Ltd")),
    ("Yokogawa", None),
    ("Refer Annexure 1", None),
    (None, None),
])
def test_parse_column_reference(value, expected):
    assert parse_column_reference(value) == expected


def _workbook(manufacturer="Refer column 12A"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MAIN"
    ws["H1"] = "MANUFACTURER:"
    ws["I1"] = manufacturer
    # header row 2, column-number row 3, data from row 4
    headers = ["ITEM NUMBER", "DESCRIPTION", "PART NUMBER", "SUPPLIER/OCM NAME", "UNIT PRICE"]
    codes = ["6", "8", "10A", "12A", "17"]
    for c, (h, code) in enumerate(zip(headers, codes), start=1):
        ws.cell(2, c, h)
        ws.cell(3, c, code)
    ws.cell(4, 1, 1); ws.cell(4, 2, "BEARING"); ws.cell(4, 4, "P&F"); ws.cell(4, 5, 10)
    profile = SheetProfile(name="MAIN")
    profile.header_row = 2
    profile.column_map = {
        "item_number": 1, "description": 2, "part_number": 3, "supplier": 4, "unit_price": 5,
    }
    return wb, profile


def test_find_column_codes_detects_number_row():
    wb, profile = _workbook()
    codes = find_column_codes(wb["MAIN"], profile.column_map)
    assert codes["12A"] == "supplier_name"
    assert codes["10A"] == "mfr_part_no"


def test_find_column_codes_none_without_number_row():
    wb, profile = _workbook()
    for c in range(1, 6):
        wb["MAIN"].cell(3, c).value = None
    assert find_column_codes(wb["MAIN"], profile.column_map) == {}


def _rows(ref):
    def spare(tag, item, supplier):
        return {"sheet": "MAIN", "tag_no": tag, "item_num": item,
                "manufacturer": ref, "supplier_name": supplier}
    return [
        {"sheet": "MAIN", "tag_no": "T1", "manufacturer": ref, "supplier_name": "YOKOGAWA"},
        spare("T1", 1, "P&F"),
        spare("T1", 2, "PROTECTRON"),
        spare("T1", 3, "P&F"),
        {"sheet": "MAIN", "tag_no": "T2", "manufacturer": ref},
        spare("T2", 4, None),
    ]


def test_spare_rows_take_referenced_column_and_equipment_most_common():
    wb, profile = _workbook()
    rows = _rows("Refer column 12A")
    resolve_column_references(wb, rows, [profile])
    assert [r["manufacturer"] for r in rows[1:4]] == ["P&F", "PROTECTRON", "P&F"]
    assert rows[0]["manufacturer"] == "P&F"          # most common of T1 spares
    assert rows[5]["manufacturer"] is None           # referenced cell blank → never the ref text
    assert rows[4]["manufacturer"] == "Refer column 12A"  # nothing to resolve from → unchanged


def test_named_value_wins_for_equipment_row():
    wb, profile = _workbook("Refer column 12A/YOKOGAWA")
    rows = _rows("Refer column 12A/YOKOGAWA")
    resolve_column_references(wb, rows, [profile])
    assert rows[0]["manufacturer"] == "YOKOGAWA"
    assert rows[4]["manufacturer"] == "YOKOGAWA"
    assert rows[1]["manufacturer"] == "P&F"          # spares still per item
    assert rows[5]["manufacturer"] == "YOKOGAWA"     # blank cell → named value


def test_sheet_own_reference_overrides_propagated_value():
    wb, profile = _workbook("Refer column 12A/YOKOGAWA")
    rows = _rows("Refer column 12A")  # propagated value lacks the named part
    resolve_column_references(wb, rows, [profile])
    assert rows[0]["manufacturer"] == "YOKOGAWA"


def test_plain_values_untouched():
    wb, profile = _workbook("Yokogawa")
    rows = _rows("YOKOGAWA")
    before = [dict(r) for r in rows]
    resolve_column_references(wb, rows, [profile])
    assert rows == before
