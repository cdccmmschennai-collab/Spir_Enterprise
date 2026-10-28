"""
Legacy .xls support: validation, the xlrd → openpyxl adapter, the BIFF8
form-control checkbox parser, SPIR-type Phase 0a, end-to-end extraction of the
two vendor .xls files, and the Celery no-retry rule for invalid files.
"""
from __future__ import annotations

import datetime as dt
import struct
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import openpyxl
import pytest
import xlrd

from spir_dynamic.analysis.header_detector import _spir_type_from_checkboxes
from spir_dynamic.extraction import workbook_loader
from spir_dynamic.extraction.file_validator import ValidationError, validate_extension, validate_file
from spir_dynamic.extraction.workbook_loader import load_workbook_any
from spir_dynamic.extraction.xls_form_controls import FormCheckbox, read_xls_checkboxes

REPO = Path(__file__).resolve().parents[1]
VENDOR_DIR = REPO / "templates" / "look for it"
XLS_0001 = VENDOR_DIR / "VEN-4499-MGDV-4-43-0001-2.xls"
XLS_0002 = VENDOR_DIR / "VEN-4499-MGDV-4-43-0002-A.xls"
_needs_vendor = pytest.mark.skipif(
    not (XLS_0001.exists() and XLS_0002.exists()), reason="vendor .xls samples not present"
)

OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
ZIP = b"PK\x03\x04"


# ── Validation ────────────────────────────────────────────────────────────────

class TestValidation:

    @pytest.mark.parametrize("name", ["a.xlsx", "a.XLSM", "a.xls"])
    def test_accepted_extensions(self, name):
        validate_extension(name)

    @pytest.mark.parametrize("name", ["a.csv", "a.xlsb", "a.ods", "a.xltx", "a.xltm", "noext", ""])
    def test_rejected_extensions(self, name):
        with pytest.raises(ValidationError):
            validate_extension(name)

    def test_csv_message_is_actionable(self):
        with pytest.raises(ValidationError, match="CSV files are not supported"):
            validate_extension("data.csv")

    @pytest.mark.parametrize("as_path", [True, False])
    @pytest.mark.parametrize("name, head, ok, match", [
        ("a.xls", OLE, True, None),
        ("a.xlsx", ZIP, True, None),
        ("a.xls", ZIP, False, "contains a ZIP-based"),          # .xlsx renamed to .xls
        ("a.xlsx", OLE, False, "not a ZIP-based"),              # .xls renamed / encrypted xlsx
        ("a.xlsm", b"garbage!", False, "bad ZIP header"),
        ("a.xls", b"garbage!", False, "bad compound-file header"),
    ])
    def test_signature_must_match_extension(self, tmp_path, as_path, name, head, ok, match):
        body = head + b"\x00" * 600
        if as_path:
            p = tmp_path / name
            p.write_bytes(body)
            content = p
        else:
            content = body
        if ok and as_path:
            validate_file(name, content)          # path mode: signature only
        elif ok:
            with pytest.raises(ValidationError, match="Cannot"):  # bytes mode also opens it
                validate_file(name, content)
        else:
            with pytest.raises(ValidationError, match=match):
                validate_file(name, content)

    def test_empty_file_rejected(self, tmp_path):
        p = tmp_path / "e.xls"
        p.write_bytes(b"")
        with pytest.raises(ValidationError, match="empty"):
            validate_file("e.xls", p)
        with pytest.raises(ValidationError, match="empty"):
            validate_file("e.xls", b"")

    @_needs_vendor
    def test_real_xls_passes_path_and_bytes(self):
        validate_file(XLS_0001.name, XLS_0001)
        validate_file(XLS_0001.name, XLS_0001.read_bytes())

    def test_loader_rejects_unsupported_extension(self):
        with pytest.raises(ValidationError, match="Unsupported file type '.xlsb'"):
            load_workbook_any(b"x", "a.xlsb")


# ── Adapter (xlrd book → openpyxl Workbook) ───────────────────────────────────

def _cell(ctype, value, xf=0):
    return SimpleNamespace(ctype=ctype, value=value, xf_index=xf)


class _FakeSheet:
    def __init__(self, name, rows, merged=(), visibility=0):
        self.name, self._rows, self.merged_cells, self.visibility = name, rows, list(merged), visibility
        self.nrows = len(rows)

    def row(self, r):
        return self._rows[r]


def _fake_book(sheets, formats):
    """formats: xf index → number-format string."""
    return SimpleNamespace(
        datemode=0,
        sheets=lambda: sheets,
        xf_list=[SimpleNamespace(format_key=i) for i in range(len(formats))],
        format_map={i: SimpleNamespace(format_str=f) for i, f in enumerate(formats)},
        release_resources=lambda: None,
    )


@pytest.fixture
def adapted():
    T, N, D, B, E, BL = (xlrd.XL_CELL_TEXT, xlrd.XL_CELL_NUMBER, xlrd.XL_CELL_DATE,
                         xlrd.XL_CELL_BOOLEAN, xlrd.XL_CELL_ERROR, xlrd.XL_CELL_BLANK)
    formats = ["General", "dd/mm/yyyy", "[h]:mm", "0.00"]
    data = _FakeSheet("Main Sheet", [
        [_cell(T, "ITEM"), _cell(N, 13.0), _cell(N, 2.5), _cell(N, 3070852610.0)],
        [_cell(D, 45658.0, 1), _cell(N, 1.5, 2), _cell(B, 1), _cell(E, 0x2A)],
        [_cell(N, 7.0, 3), _cell(BL, ""), _cell(xlrd.XL_CELL_EMPTY, ""), _cell(N, 1e20)],
        [_cell(T, "merged"), _cell(BL, ""), _cell(BL, ""), _cell(BL, "")],
    ], merged=[(3, 4, 0, 3)])
    hidden = _FakeSheet("Annexure-1", [[_cell(T, "x")]], visibility=1)
    very = _FakeSheet("Calc", [[_cell(T, "y")]], visibility=2)
    empty = _FakeSheet("Empty", [])
    book = _fake_book([data, hidden, very, empty], formats)
    with patch.object(workbook_loader.xlrd, "open_workbook", return_value=book), \
         patch.object(workbook_loader, "_form_checkboxes", return_value=None):
        return workbook_loader.load_xls_bytes(b"ignored")


class TestAdapter:

    def test_sheet_names_and_states(self, adapted):
        assert adapted.sheetnames == ["Main Sheet", "Annexure-1", "Calc", "Empty"]
        assert [ws.sheet_state for ws in adapted.worksheets] == ["visible", "hidden", "veryHidden", "visible"]

    def test_numbers_normalised_like_openpyxl(self, adapted):
        ws = adapted["Main Sheet"]
        assert ws.cell(1, 2).value == 13 and type(ws.cell(1, 2).value) is int
        assert ws.cell(1, 3).value == 2.5
        assert ws.cell(1, 4).value == 3070852610 and type(ws.cell(1, 4).value) is int
        assert ws.cell(3, 1).value == 7 and type(ws.cell(3, 1).value) is int   # "0.00" is not a date
        assert type(ws.cell(3, 4).value) is float                               # beyond int limit

    def test_dates_durations_booleans_errors(self, adapted):
        ws = adapted["Main Sheet"]
        assert ws.cell(2, 1).value == dt.datetime(2025, 1, 1)
        assert ws.cell(2, 2).value == dt.timedelta(hours=36)
        assert ws.cell(2, 3).value is True
        assert ws.cell(2, 4).value == "#N/A"

    def test_blank_cells_keep_dimensions(self, adapted):
        ws = adapted["Main Sheet"]
        assert ws.cell(3, 2).value is None
        assert (ws.max_row, ws.max_column) == (4, 4)

    def test_merged_ranges(self, adapted):
        ws = adapted["Main Sheet"]
        assert [str(r) for r in ws.merged_cells.ranges] == ["A4:C4"]
        assert ws.cell(4, 1).value == "merged"
        assert isinstance(ws.cell(4, 2), openpyxl.cell.cell.MergedCell)

    def test_empty_sheet(self, adapted):
        ws = adapted["Empty"]
        assert (ws.max_row, ws.max_column) == (1, 1) and ws.cell(1, 1).value is None

    def test_no_zip_handles_for_xls(self, adapted):
        assert adapted._spir_raw_bytes is None and adapted._spir_raw_path is None

    def test_encrypted_xls_is_validation_error(self):
        with patch.object(workbook_loader.xlrd, "open_workbook",
                          side_effect=xlrd.XLRDError("Workbook is encrypted")):
            with pytest.raises(ValidationError, match="password-protected"):
                workbook_loader.load_xls_bytes(b"x")


# ── BIFF8 checkbox parser ─────────────────────────────────────────────────────

def _rec(rid, body=b""):
    return struct.pack("<HH", rid, len(body)) + body


def _checkbox_obj(checked=0, with_data=True, overrun=False):
    cmo = struct.pack("<HHHHH", 0x15, 0x12, 0x0B, 1, 0) + b"\x00" * 12
    subs = struct.pack("<HH", 0x0A, 12) + b"\x00" * 12
    if with_data:
        subs += struct.pack("<HHH", 0x12, 8, checked) + b"\x00" * 6
    if overrun:
        subs += struct.pack("<HH", 0x13, 200)
    return _rec(0x5D, cmo + subs + struct.pack("<HH", 0, 0))


def _picture_obj():
    return _rec(0x5D, struct.pack("<HHHHH", 0x15, 0x12, 0x08, 2, 0) + b"\x00" * 12 + b"\x00" * 4)


def _caption(text):
    txo = b"\x00" * 10 + struct.pack("<HH", len(text), 16) + b"\x00" * 4
    return _rec(0x1B6, txo) + _rec(0x3C, b"\x00" + text.encode("latin-1")) + _rec(0x3C, b"\x00" * 16)


def _biff_stream(*sheet_bodies, encrypted=False):
    """Globals (BOF, [FILEPASS], BOUNDSHEETs, EOF) followed by one substream per sheet."""
    bof = _rec(0x0809, struct.pack("<HH", 0x0600, 0x0005) + b"\x00" * 12)
    sheet_bof = _rec(0x0809, struct.pack("<HH", 0x0600, 0x0010) + b"\x00" * 12)
    eof = _rec(0x000A)

    def globals_(offsets):
        out = bof + (_rec(0x2F, b"\x00" * 4) if encrypted else b"")
        for i, off in enumerate(offsets):
            name = f"S{i}".encode()
            out += _rec(0x85, struct.pack("<IBB", off, 0, 0) + bytes([len(name), 0]) + name)
        return out + eof

    size = len(globals_([0] * len(sheet_bodies)))
    offsets, pos = [], size
    for body in sheet_bodies:
        offsets.append(pos)
        pos += len(sheet_bof) + len(body) + len(eof)
    return globals_(offsets) + b"".join(sheet_bof + b + eof for b in sheet_bodies)


def _four(states):
    names = ["Commissioning Spares", "Initial Spares", "Normal Operating Spares", "Life Cycle Spares"]
    return b"".join(_rec(0xEC, b"\x00" * 8) + _checkbox_obj(s) + _rec(0xEC, b"\x00" * 8) + _caption(n)
                    for n, s in zip(names, states))


class TestCheckboxParser:

    def test_reads_caption_and_state_per_sheet(self):
        boxes = read_xls_checkboxes(_biff_stream(_four([0, 1, 0, 0]), _picture_obj()))
        assert boxes[0] == [
            FormCheckbox("Commissioning Spares", 0), FormCheckbox("Initial Spares", 1),
            FormCheckbox("Normal Operating Spares", 0), FormCheckbox("Life Cycle Spares", 0),
        ]
        assert boxes[1] == []          # pictures are not checkboxes

    def test_missing_state_and_overrun_are_malformed(self):
        body = _checkbox_obj(with_data=False) + _caption("Initial Spares") + \
               _checkbox_obj(1, overrun=True) + _caption("Life Cycle Spares")
        assert read_xls_checkboxes(_biff_stream(body))[0] == [
            FormCheckbox("Initial Spares", None), FormCheckbox("Life Cycle Spares", None),
        ]

    def test_checkbox_without_caption(self):
        assert read_xls_checkboxes(_biff_stream(_checkbox_obj(1)))[0] == [FormCheckbox(None, 1)]

    def test_encrypted_or_not_biff8_gives_none(self):
        assert read_xls_checkboxes(_biff_stream(_four([0, 1, 0, 0]), encrypted=True)) is None
        assert read_xls_checkboxes(b"\x00" * 40) is None
        assert read_xls_checkboxes(b"") is None


class TestSpirTypeFromCheckboxes:

    @staticmethod
    def _boxes(states, names=None):
        names = names or ["Commissioning Spares", "Initial Spares", "Normal Operating Spares", "Life Cycle Spares"]
        return [FormCheckbox(n, s) for n, s in zip(names, states)]

    def test_single_selection(self):
        assert _spir_type_from_checkboxes(self._boxes([0, 1, 0, 0])) == (True, "initial spares")
        assert _spir_type_from_checkboxes(self._boxes([0, 0, 1, 0])) == (True, "normal operating spares")

    def test_no_selection_is_null(self):
        assert _spir_type_from_checkboxes(self._boxes([0, 0, 0, 0])) == (True, None)

    def test_conflicting_selection_is_null(self):
        assert _spir_type_from_checkboxes(self._boxes([1, 1, 0, 0])) == (True, None)

    def test_malformed_state_is_null(self):
        assert _spir_type_from_checkboxes(self._boxes([0, 1, None, 0])) == (True, None)
        assert _spir_type_from_checkboxes(self._boxes([0, 1, 2, 0])) == (True, None)   # "mixed"

    def test_ambiguous_caption_is_null(self):
        boxes = [FormCheckbox("Initial Spares / Normal Operating Spares", 1), FormCheckbox("Life Cycle Spares", 0)]
        assert _spir_type_from_checkboxes(boxes) == (True, None)

    def test_uncaptioned_box_alongside_types_is_null(self):
        assert _spir_type_from_checkboxes(self._boxes([0, 1]) + [FormCheckbox(None, 0)]) == (True, None)

    def test_unrelated_controls_fall_through(self):
        assert _spir_type_from_checkboxes([FormCheckbox("Approved", 1), FormCheckbox(None, 0)]) == (False, None)
        assert _spir_type_from_checkboxes([]) == (False, None)


# ── End-to-end on the vendor files ────────────────────────────────────────────

def _run_pipeline(source, name):
    from spir_dynamic.app import pipeline
    from spir_dynamic.services.planning_plant import PlanningPlant

    store = SimpleNamespace(put=lambda *a, **k: None)
    with patch.object(pipeline, "_apply_currency_conversion", return_value=None), \
         patch.object(pipeline, "get_storage", return_value=store):
        return pipeline.run_pipeline(source, name, PlanningPlant("2400", "NGL Mesaieed"))


@_needs_vendor
class TestVendorXls:

    @pytest.mark.parametrize("path, spir_type, spir_no", [
        (XLS_0001, "initial spares", "VEN-4499-MGDV-4-43-0001-REV-2"),
        (XLS_0002, "normal operating spares", "VEN-4499-MGDV-4-43-0002-REV-A"),
    ])
    @pytest.mark.parametrize("as_bytes", [False, True])
    def test_extraction(self, path, spir_type, spir_no, as_bytes):
        res = _run_pipeline(path.read_bytes() if as_bytes else path, path.name)
        assert res["spir_type"] == spir_type
        assert res["spir_no"] == spir_no
        cols = res["preview_cols"]
        tag, item = cols.index("TAG NO"), cols.index("ITEM NUMBER")
        rows = res["preview_rows"]
        assert {r[tag] for r in rows} == {"V-ESDV-754", "V-ESDV-755"}
        assert not any("ANNEX" in str(r[tag]).upper() for r in rows)
        assert all(r[cols.index("SPIR TYPE")] == spir_type.upper() for r in rows)
        # Every source spare item appears exactly once per tag.
        expected_items = {XLS_0001: 35, XLS_0002: 39}[path]
        for t in ("V-ESDV-754", "V-ESDV-755"):
            items = [r[item] for r in rows if r[tag] == t and r[item] not in (None, "")]
            assert sorted(items) == list(range(1, expected_items + 1))

    def test_source_fields_carried_through(self):
        res = _run_pipeline(XLS_0001, XLS_0001.name)
        cols = res["preview_cols"]
        c = {k: cols.index(k) for k in ("ITEM NUMBER", "TAG NO", "DESCRIPTION OF PARTS",
                                         "MANUFACTURER PART NUMBER", "UNIT PRICE", "UNIT OF MEASURE")}
        item1 = next(r for r in res["preview_rows"] if r[c["TAG NO"]] == "V-ESDV-754" and r[c["ITEM NUMBER"]] == 1)
        assert item1[c["DESCRIPTION OF PARTS"]].startswith("SEAT RING + DELTA RING/INSERT")
        assert item1[c["MANUFACTURER PART NUMBER"]] == "ORD-22343-1,2-SR+INSERT"
        assert item1[c["UNIT PRICE"]] == 3200
        assert item1[c["UNIT OF MEASURE"]] == "SET"

    def test_vendor_checkboxes_and_misdirected_link(self):
        wb = load_workbook_any(XLS_0002, XLS_0002.name)
        boxes = wb._spir_form_checkboxes["SPIR"]
        assert {b.caption: b.checked for b in boxes} == {
            "Commissioning Spares": 0, "Normal Operating Spares": 1,
            "Initial Spares": 0, "Life Cycle Spares": 0,
        }
        # The Initial Spares box is linked to U3, which holds manufacturer text;
        # it must not register as a selection.
        assert "JC" in str(wb["SPIR"]["U3"].value)


# ── Celery: deterministic file errors are not retried ────────────────────────

class TestNoRetryForInvalidFiles:

    def _run(self, tmp_path, side_effect):
        import importlib.util
        from spir_dynamic.tasks.extraction_tasks import process_file_task

        # Reuse the task harness of test_delivery_cap (tests/ is not a package).
        spec = importlib.util.spec_from_file_location("_delivery_cap", Path(__file__).with_name("test_delivery_cap.py"))
        harness = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(harness)
        _PLANT_FIELDS, _task_context = harness._PLANT_FIELDS, harness._task_context

        upload = tmp_path / "bad.xls"
        upload.write_bytes(OLE + b"\x00" * 100)
        with patch.object(process_file_task, "retry", side_effect=RuntimeError("retried")) as retry, \
             _task_context(str(upload), 1, extra_patches=[
                 patch("spir_dynamic.app.pipeline.run_pipeline", side_effect=side_effect)]) as (task, *_):
            try:
                result = task.run(**_PLANT_FIELDS, job_id="job-1", file_idx=0,
                                  source_key=upload.name, filename="bad.xls", user_id="u")
            except RuntimeError:
                result = None
        return result, retry

    def test_validation_error_fails_without_retry(self, tmp_path):
        result, retry = self._run(tmp_path, ValidationError("This .xls workbook is password-protected."))
        assert result["status"] == "error" and "password-protected" in result["error"]
        retry.assert_not_called()

    def test_transient_error_still_retries(self, tmp_path):
        _, retry = self._run(tmp_path, ConnectionError("storage blip"))
        retry.assert_called_once()
