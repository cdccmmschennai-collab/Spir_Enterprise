"""
Saipem SPIR support (extraction/saipem).

Synthetic workbooks (tests/saipem_fixtures.py) cover structure and variants;
the customer's workbook is used only when present locally (it is not in
version control) and those tests skip otherwise.
"""
from __future__ import annotations

import io
from pathlib import Path
from unittest.mock import MagicMock, patch

import openpyxl
import pytest

from spir_dynamic.extraction.output_schema import OUTPUT_COLS
from spir_dynamic.extraction.saipem import detect_saipem_sheets
from spir_dynamic.extraction.saipem.equipment import expand_equipment_tag
from spir_dynamic.extraction.saipem.excel_builder import build_saipem_xlsx
from spir_dynamic.extraction.saipem.extractor import extract_saipem
from spir_dynamic.extraction.saipem.layout import resolve_layout
from spir_dynamic.extraction.saipem.matrix import vendor_total
from spir_dynamic.extraction.saipem.profile import SAIPEM_COLUMNS
from spir_dynamic.extraction.saipem.values import identifier_text
from spir_dynamic.services.planning_plant import PlanningPlant
from tests.saipem_fixtures import (
    CAPITAL_FIELDS,
    OPERATING_FIELDS,
    SPARE_FIELDS,
    TAIL_FIELDS,
    Equipment,
    Field,
    Line,
    saipem_workbook,
    two_pump_lines,
)

REPO_ROOT = Path(__file__).resolve().parents[1]
REAL_XLS = REPO_ROOT / "EX2515011 - FINAL SPIR ( REV 5).xls"
needs_real_xls = pytest.mark.skipif(not REAL_XLS.exists(), reason="customer workbook not present")
GENERIC_WORKBOOK = REPO_ROOT / "templates" / "inputs" / "VEN-4460-DGTYP-5-43-1202-A.xlsm"
NGL = PlanningPlant("2400", "NGL Mesaieed")

PUMP_A, PUMP_B = "005400PH011A", "005400PH011B"
PA_A, PA_B = "004300PA001A", "004300PA001B"


def _extract(wb=None):
    wb = wb or saipem_workbook()
    return extract_saipem(wb, detect_saipem_sheets(wb))


def _rows_for(result, tag, spares_only=True):
    return [r for r in result.rows
            if r["item_tag_no"] == tag and (not spares_only or r.get("spir_pos_no"))]


def _by_pos(result, tag):
    return {r["spir_pos_no"]: r for r in _rows_for(result, tag)}


def _roundtrip(wb) -> openpyxl.Workbook:
    buf = io.BytesIO()
    wb.save(buf)
    return openpyxl.load_workbook(io.BytesIO(buf.getvalue()), data_only=True)


def _run_pipeline(path: Path, filename: str):
    """run_pipeline with storage captured; generic-only steps must not run."""
    from spir_dynamic.app.pipeline import run_pipeline
    stored: dict = {}
    storage = MagicMock()
    storage.put.side_effect = lambda fid, data, name, **kw: stored.update(xlsx=data, name=name)
    boom = AssertionError("generic step must not run for Saipem")
    with patch("spir_dynamic.extraction.saipem.service.get_storage", return_value=storage), \
         patch("spir_dynamic.app.pipeline.get_storage", return_value=storage), \
         patch("spir_dynamic.app.pipeline.extract_workbook", side_effect=boom), \
         patch("spir_dynamic.app.pipeline._apply_plant", side_effect=boom), \
         patch("spir_dynamic.app.pipeline._apply_currency_conversion", side_effect=boom), \
         patch("spir_dynamic.app.pipeline.post_process_rows", side_effect=boom), \
         patch("spir_dynamic.app.pipeline.deduplicate_rows", side_effect=boom), \
         patch("spir_dynamic.app.pipeline.build_xlsx", side_effect=boom):
        result = run_pipeline(path, filename, NGL)
    return result, openpyxl.load_workbook(io.BytesIO(stored["xlsx"]))


EXPECTED_HEADER = [
    # (band, field, sub) per column, in export order
    (None, "SPIR No.", None),
    ("EQUIPMENT IDENTIFICATION", "ITEM TAG No. (MAIN ITEM)", None),
    ("EQUIPMENT IDENTIFICATION", "Free Field (Second TAG Codification or Item Criticality)", None),
    ("EQUIPMENT IDENTIFICATION", "SERIAL No.", None),
    ("EQUIPMENT IDENTIFICATION", "MODEL OR TYPE", None),
    ("EQUIPMENT IDENTIFICATION", "EQUIPMENT SUB ITEM No.", None),
    (None, "No. OF UNITS", None),
    (None, "QUANTITY OF PARTS INSTALLED PER ONE UNIT OF EQUIPMENT SHOWN IN COLUMN A-E", None),
    (None, "TOTAL QUANTITY OF IDENTICAL PARTS INSTALLED", None),
    ("SPARE PARTS IDENTIFICATION", "SPIR POS. No. (from 0001)", None),
    ("SPARE PARTS IDENTIFICATION", "DESCRIPTION OF PARTS", None),
    ("SPARE PARTS IDENTIFICATION", "NEW DESCRIPTION OF PARTS", None),
    ("SPARE PARTS IDENTIFICATION", "MATERIAL", None),
    ("SPARE PARTS IDENTIFICATION", "ORIGINAL MAKERS NAME", None),
    ("SPARE PARTS IDENTIFICATION", "ORIGINAL MAKERS PART No.", None),
    ("SPARE PARTS IDENTIFICATION", "RELEVANT DOCUMENT No.", "DRAWING No."),
    ("SPARE PARTS IDENTIFICATION", "RELEVANT DOCUMENT No.", "PAGE"),
    ("SPARE PARTS IDENTIFICATION", "RELEVANT DOCUMENT No.", "POS. No."),
    ("SPARE PARTS IDENTIFICATION", "VENDOR CODE No.", None),
    ("SPARE PARTS IDENTIFICATION", "Business unit codification (Miap code)", None),
    ("QUANTITY FOR 2 YEARS OPERATION", "UNIT of MEAS.", None),
    ("QUANTITY FOR 2 YEARS OPERATION", "RECOMMEND. By VENDOR", "Total quantity"),
    ("QUANTITY FOR 2 YEARS OPERATION", "RECOMMEND. By VENDOR", "Quantity for PM"),
    ("QUANTITY FOR 2 YEARS OPERATION", "RECOMMEND. By VENDOR", "Quantity for CM"),
    ("QUANTITY FOR 2 YEARS OPERATION", "ENGINEER. RECOMMEND.", None),
    ("QUANTITY FOR 2 YEARS OPERATION", "OPERATION RECOMMEND.", None),
    ("QUANTITY FOR 2 YEARS OPERATION", "PURCHASED QUANTITY (approved)", None),
    ("CAPITAL SPARES", "RECOMMEND. By VENDOR", None),
    ("CAPITAL SPARES", "ENGINEER. RECOMMEND.", None),
    ("CAPITAL SPARES", "OPERATION RECOMMEND.", None),
    ("CAPITAL SPARES", "PURCHASED QUANTITY (approved)", None),
    (None, "WEIGHT", None),
    (None, "CHARACTERISTIC DIMENSION", None),
    (None, "OBSOLESCENCE", None),
    (None, "G/S", None),
    (None, "SHELF LIFE", None),
    (None, "PRICE", "UNIT"),
    (None, "PRICE", "AMOUNT"),
    (None, "PRICE", "CURRENCY"),
    (None, "DELIVERY TIME (IN WEEKS)", None),
    (None, "VENDOR'S NAME", None),
    (None, "ADDRESS", None),
    (None, "PHONE No. / FAX No.", None),
    (None, "E-MAIL ADDRESS", None),
]


def _header_triples(ws) -> list[tuple]:
    """(band, field, sub) per output column, resolving merged header cells."""
    def value(r, c):
        for rng in ws.merged_cells.ranges:
            if rng.min_row <= r <= rng.max_row and rng.min_col <= c <= rng.max_col:
                return ws.cell(rng.min_row, rng.min_col).value, (rng.min_row, rng.max_row)
        return ws.cell(r, c).value, (r, r)

    out = []
    for c in range(1, ws.max_column + 1):
        v1, span1 = value(1, c)
        v2, span2 = value(2, c)
        v3, span3 = value(3, c)
        if span1 == (1, 3):            # field merged down rows 1-3
            out.append((None, v1, None))
        elif span1 == (1, 2):          # field over rows 1-2, sub-header on row 3
            out.append((None, v1, v3))
        elif span2 == (2, 3):          # band on row 1, field merged rows 2-3
            out.append((v1, v2, None))
        else:                          # band, field, sub-header
            out.append((v1, v2, v3))
    return out


# ── 1-2. Detection and sheet exclusion ──────────────────────────────────────

class TestDetection:
    def test_form_detected_example_and_instruction_excluded(self):
        assert detect_saipem_sheets(_roundtrip(saipem_workbook())) == ["SPIR Form"]

    @pytest.mark.parametrize("name", ["SPIR Form Example", "Sample SPIR", "EXAMPLE", "spir form (sample)"])
    def test_example_or_sample_name_excluded_although_structurally_valid(self, name):
        wb = saipem_workbook(with_example=False, with_instruction=False)
        wb["SPIR Form"].title = name
        assert detect_saipem_sheets(wb) == []

    def test_instruction_sheet_rejected_by_structure_even_when_renamed(self):
        wb = saipem_workbook(with_example=False)
        wb["Instruction"].title = "Notes"
        assert detect_saipem_sheets(wb) == ["SPIR Form"]

    def test_example_values_never_reach_output(self):
        result = _extract(_roundtrip(saipem_workbook()))
        tags = {r["item_tag_no"] for r in result.rows}
        assert tags == {PUMP_A, PUMP_B, PA_A, PA_B}
        assert not any("XX" in str(r.get("serial_no")) for r in result.rows)
        assert result.sheets == ["SPIR Form"]

    def test_detection_does_not_depend_on_filename_or_sheet_name(self):
        wb = saipem_workbook(with_example=False, with_instruction=False)
        wb["SPIR Form"].title = "Sheet1"
        assert detect_saipem_sheets(wb) == ["Sheet1"]

    def test_hidden_sheet_skipped(self):
        wb = saipem_workbook(with_example=False, with_instruction=False)
        wb.create_sheet("visible")
        wb["SPIR Form"].sheet_state = "hidden"
        assert detect_saipem_sheets(wb) == []

    def test_generic_workbook_not_claimed_and_not_modified(self):
        if not GENERIC_WORKBOOK.exists():
            pytest.skip("generic template not present")
        wb = openpyxl.load_workbook(GENERIC_WORKBOOK, data_only=True, keep_links=False)
        dims = {ws.title: (ws.max_row, ws.max_column) for ws in wb.worksheets}
        assert detect_saipem_sheets(wb) == []
        assert dims == {ws.title: (ws.max_row, ws.max_column) for ws in wb.worksheets}

    def test_plain_table_not_claimed(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["TAG NO", "DESCRIPTION OF PARTS", "SPIR POS. No.", "SERIAL No.", "MODEL OR TYPE"])
        ws.append(["P-101", "SEAL", "0001", "S1", "M1"])
        assert detect_saipem_sheets(wb) == []


# ── 3, 4. Reading .xlsx / .xls through the pipeline; revision from the sheet ─

class TestPipeline:
    def test_xlsx_through_run_pipeline(self, tmp_path):
        path = tmp_path / "EX2515011 - FINAL SPIR ( REV 5).xlsx"
        saipem_workbook().save(path)
        result, out = _run_pipeline(path, path.name)
        assert result["format"] == "SAIPEM: SPIR Form"
        assert result["spir_no"] == "102700DKEZFW657-REV-4"      # sheet REV 4, not filename REV 5
        assert result["total_tags"] == 4 and result["spare_items"] == 3
        assert result["total_rows"] == len(result["preview_rows"]) == out.active.max_row - 3
        assert result["spir_type"] is None and result["currency_rates"] is None
        assert result["warnings"] == []

    @needs_real_xls
    def test_real_xls_through_run_pipeline(self, tmp_path):
        import shutil
        path = tmp_path / REAL_XLS.name
        shutil.copy(REAL_XLS, path)
        result, out = _run_pipeline(path, path.name)
        assert result["spir_no"] == "102700DKEZFW657-REV-4"
        assert result["total_tags"] == 4 and result["spare_items"] == 47
        assert result["total_rows"] == 98            # 4 equipment rows + 2×33 + 2×14 spare rows
        assert result["warnings"] == []
        assert _header_triples(out.active) == EXPECTED_HEADER

    def test_revision_falls_back_to_revision_block(self):
        wb = saipem_workbook(rev=None, revision_block=7)
        assert _extract(wb).spir_no == "102700DKEZFW657-REV-7"

    def test_no_revision_keeps_plain_spir_no(self):
        assert _extract(saipem_workbook(rev=None)).spir_no == "102700DKEZFW657"

    def test_response_keys_match_generic_pipeline(self, tmp_path):
        path = tmp_path / "form.xlsx"
        saipem_workbook().save(path)
        result, _ = _run_pipeline(path, path.name)
        generic_keys = {"status", "file_id", "filename", "format", "spir_no", "equipment",
                        "manufacturer", "supplier", "spir_type", "eqpt_qty", "spare_items",
                        "total_tags", "annexure_count", "total_rows", "dup1_count", "sap_count",
                        "preview_cols", "preview_rows", "sheet_profiles", "currency_rates",
                        "plant", "plant_description"}
        assert generic_keys <= set(result)


# ── 5. Exact headers, order, repeated groups ────────────────────────────────

class TestOutputLayout:
    def test_exact_ordered_three_level_header(self):
        result = _extract()
        ws = openpyxl.load_workbook(io.BytesIO(
            build_saipem_xlsx(result.rows, result.columns, result.band_labels, result.spir_no))).active
        assert _header_triples(ws) == EXPECTED_HEADER

    def test_repeated_labels_stay_distinct(self):
        ids = [c.id for c in SAIPEM_COLUMNS]
        assert len(ids) == len(set(ids))
        from spir_dynamic.extraction.saipem.profile import DEFAULT_BAND_LABELS, flat_label
        flat = [flat_label(c, DEFAULT_BAND_LABELS) for c in SAIPEM_COLUMNS]
        assert len(flat) == len(set(flat))
        assert "QUANTITY FOR 2 YEARS OPERATION | PURCHASED QUANTITY (approved)" in flat
        assert "CAPITAL SPARES | PURCHASED QUANTITY (approved)" in flat

    def test_values_land_under_their_source_label(self):
        result = _extract()
        ws = openpyxl.load_workbook(io.BytesIO(
            build_saipem_xlsx(result.rows, result.columns, result.band_labels, result.spir_no))).active
        header = _header_triples(ws)
        rows = list(ws.iter_rows(min_row=4, values_only=True))
        line2 = next(r for r in rows if r[header.index(EXPECTED_HEADER[9])] == "0002")
        get = lambda triple: line2[header.index(triple)]
        assert get(("SPARE PARTS IDENTIFICATION", "DESCRIPTION OF PARTS", None)) == "KEY - LINE SHAFT"
        assert get(("SPARE PARTS IDENTIFICATION", "RELEVANT DOCUMENT No.", "POS. No.")) == "2120.1"
        assert get(("QUANTITY FOR 2 YEARS OPERATION", "RECOMMEND. By VENDOR", "Quantity for PM")) == 5
        assert get(("QUANTITY FOR 2 YEARS OPERATION", "RECOMMEND. By VENDOR", "Quantity for CM")) == 5
        assert get(("QUANTITY FOR 2 YEARS OPERATION", "RECOMMEND. By VENDOR", "Total quantity")) == 10
        assert get(("CAPITAL SPARES", "RECOMMEND. By VENDOR", None)) == 0
        assert get((None, "QUANTITY OF PARTS INSTALLED PER ONE UNIT OF EQUIPMENT SHOWN IN COLUMN A-E", None)) == 8
        assert get((None, "TOTAL QUANTITY OF IDENTICAL PARTS INSTALLED", None)) == 16
        assert get((None, "PRICE", "CURRENCY")) == "USD"
        assert ws.cell(4 + rows.index(line2), header.index(EXPECTED_HEADER[9]) + 1).number_format == "@"

    def test_data_cells_never_merged(self):
        result = _extract()
        ws = openpyxl.load_workbook(io.BytesIO(
            build_saipem_xlsx(result.rows, result.columns, result.band_labels, result.spir_no))).active
        assert all(rng.max_row <= 3 for rng in ws.merged_cells.ranges)

    def test_no_generic_only_columns(self):
        labels = {c.field for c in SAIPEM_COLUMNS} | {c.sub for c in SAIPEM_COLUMNS if c.sub}
        for generic in ("PLANNING PLANT", "PLANNING PLANT DESCRIPTION", "UNIT PRICE (QAR)",
                        "OLD MATERIAL NUMBER/SPF NUMBER", "POSITION NUMBER", "ERROR", "S.NO",
                        "SUPPLIER/ OCM NAME", "SHEET", "SPIR TYPE"):
            assert generic not in labels


# ── 6-7. Tag expansion and serial pairing ───────────────────────────────────

class TestTags:
    @pytest.mark.parametrize("raw, expected", [
        ("005400PH011A/B", [PUMP_A, PUMP_B]),
        ("004300PA001A/B", [PA_A, PA_B]),
        ("P-3425 A/B/C", ["P-3425 A", "P-3425 B", "P-3425 C"]),
        ("005400PH011A / 005400PH011B", [PUMP_A, PUMP_B]),
        ("23V01-A/B", ["23V01-A", "23V01-B"]),
        ("005400PH011A", [PUMP_A]),
        ("  005400PH011A  ", [PUMP_A]),
        (None, []),
    ])
    def test_expand(self, raw, expected):
        assert expand_equipment_tag(raw)[0] == expected

    @pytest.mark.parametrize("raw", ["PUMP/MOTOR", "MTR-EX2515011-PUMP/MOTOR-SET", "A/B", "SEE NOTE / TBA"])
    def test_non_tag_expressions_kept_whole(self, raw):
        tags, warning = expand_equipment_tag(raw)
        assert tags == [raw] and warning

    def test_both_groups_expand_with_positional_serials(self):
        result = _extract()
        serial = {r["item_tag_no"]: r["serial_no"] for r in result.rows}
        assert serial == {PUMP_A: "1825076", PUMP_B: "1825077", PA_A: "EX13251310", PA_B: "EX13251309"}

    def test_part_numbers_with_slash_never_split(self):
        result = _extract()
        dpos = {r["rel_doc_pos_no"] for r in _rows_for(result, PA_A)}
        assert dpos == {"4590.1/4590.2"}

    def test_single_group_serial_repeats_for_every_tag(self):
        wb = saipem_workbook(equipment=[Equipment("005400PH011A/B", "-", "08ELM", units=2)],
                             lines=[Line({"pos": "0001", "desc": "SEAL"}, per_unit={0: 1})])
        assert {r["serial_no"] for r in _extract(wb).rows} == {"-"}

    def test_fewer_serials_than_tags_pairs_by_position(self):
        wb = saipem_workbook(equipment=[Equipment("P-100 A/B/C", "S1/S2", "M", units=3)],
                             lines=[Line({"pos": "0001", "desc": "SEAL"}, per_unit={0: 1})])
        result = _extract(wb)
        serial = {r["item_tag_no"]: r["serial_no"] for r in result.rows}
        assert serial == {"P-100 A": "S1", "P-100 B": "S2", "P-100 C": None}
        assert any("serial" in w for w in result.warnings)

    def test_unit_mismatch_does_not_fail(self):
        wb = saipem_workbook(equipment=[Equipment("005400PH011A/B", "1/2", "08ELM", units=3)],
                             lines=[Line({"pos": "0001", "desc": "SEAL"}, per_unit={0: 1})])
        result = _extract(wb)
        assert {r["item_tag_no"] for r in result.rows} == {PUMP_A, PUMP_B}
        assert any("No. OF UNITS is 3" in w for w in result.warnings)

    def test_equipment_metadata_repeated_per_tag(self):
        result = _extract()
        for tag in (PUMP_A, PUMP_B):
            assert {r["model_or_type"] for r in result.rows if r["item_tag_no"] == tag} == {"08ELM"}
            assert {r["free_field"] for r in result.rows if r["item_tag_no"] == tag} == {"-"}


# ── 8. Equipment quantity — generic convention ──────────────────────────────

class TestEquipmentQuantity:
    def test_units_only_on_equipment_row(self):
        result = _extract()
        for tag in (PUMP_A, PUMP_B, PA_A, PA_B):
            rows = [r for r in result.rows if r["item_tag_no"] == tag]
            assert rows[0].get("spir_pos_no") is None and rows[0]["no_of_units"] == 2
            assert all(r.get("no_of_units") is None for r in rows[1:])

    def test_missing_units_left_blank_not_tag_count(self):
        wb = saipem_workbook(equipment=[Equipment("005400PH011A/B", "1/2", "08ELM", units=None)],
                             lines=[Line({"pos": "0001", "desc": "SEAL"}, per_unit={0: 1})])
        assert {r.get("no_of_units") for r in _extract(wb).rows} == {None}

    @needs_real_xls
    def test_matches_generic_eqpt_qty_on_real_file(self, tmp_path):
        from spir_dynamic.extraction.unified_extractor import extract_workbook
        from spir_dynamic.extraction.xls_converter import convert_xls_to_xlsx
        xlsx = convert_xls_to_xlsx(REAL_XLS, out_dir=tmp_path)
        wb = openpyxl.load_workbook(xlsx, data_only=True)
        wb._spir_raw_bytes, wb._spir_raw_path = None, str(xlsx)
        generic = {r["tag_no"]: r.get("eqpt_qty")
                   for r in extract_workbook(wb, REAL_XLS.name)["rows"]
                   if r.get("sheet", "").upper() == "SPIR FORM" and r.get("eqpt_qty") is not None}
        saipem = {r["item_tag_no"]: r["no_of_units"]
                  for r in _extract(openpyxl.load_workbook(xlsx, data_only=True)).rows
                  if r.get("no_of_units") is not None}
        assert generic == saipem == {PUMP_A: 2, PUMP_B: 2, PA_A: 2, PA_B: 2}


# ── 9-11. Spare association, repeated and separated quantities ──────────────

class TestSpareAssociation:
    def test_each_group_gets_only_its_lines(self):
        result = _extract()
        assert set(_by_pos(result, PUMP_A)) == set(_by_pos(result, PUMP_B)) == {"0001", "0002"}
        assert set(_by_pos(result, PA_A)) == set(_by_pos(result, PA_B)) == {"0003"}

    def test_row_convention_equipment_row_then_spares(self):
        result = _extract()
        assert [r["item_tag_no"] for r in result.rows] == (
            [PUMP_A] * 3 + [PUMP_B] * 3 + [PA_A] * 2 + [PA_B] * 2)

    def test_spare_quantities_repeated_unchanged_per_tag(self):
        result = _extract()
        a, b = _by_pos(result, PUMP_A)["0002"], _by_pos(result, PUMP_B)["0002"]
        for fid in ("op2y_vendor_pm", "op2y_vendor_cm", "op2y_vendor_total", "cap_vendor",
                    "qty_installed_per_unit", "total_qty_identical_installed"):
            assert a[fid] == b[fid]
        assert (a["op2y_vendor_pm"], a["op2y_vendor_cm"], a["op2y_vendor_total"]) == (5, 5, 10)

    def test_installed_and_spare_quantities_kept_apart(self):
        line = _by_pos(_extract(), PUMP_A)["0002"]
        assert line["qty_installed_per_unit"] == 8            # per unit, this tag's column
        assert line["total_qty_identical_installed"] == 16    # source TOTAL column
        assert line["op2y_vendor_pm"] == 5 and line["op2y_vendor_cm"] == 5
        assert line["op2y_purchased"] is None and line["cap_purchased"] is None

    def test_zero_and_text_per_unit_values_associate(self):
        lines = [Line({"pos": "0001", "desc": "A"}, per_unit={0: 0}),
                 Line({"pos": "0002", "desc": "B"}, per_unit={0: "AS REQ"}),
                 Line({"pos": "0003", "desc": "C"}, per_unit={0: "   "}),
                 Line({"pos": "0004", "desc": "D"}, per_unit={1: 1})]
        result = _extract(saipem_workbook(lines=lines))
        by_pos = _by_pos(result, PUMP_A)
        assert set(by_pos) == {"0001", "0002"}
        assert by_pos["0001"]["qty_installed_per_unit"] == 0
        assert by_pos["0002"]["qty_installed_per_unit"] == "AS REQ"
        assert any("0003" in w for w in result.warnings)       # associated with no tag → reported

    def test_tag_without_spares_skipped_like_generic(self):
        lines = [Line({"pos": "0001", "desc": "A"}, per_unit={0: 1})]
        result = _extract(saipem_workbook(lines=lines))
        assert {r["item_tag_no"] for r in result.rows} == {PUMP_A, PUMP_B}
        assert any(PA_A in w for w in result.warnings)

    @needs_real_xls
    def test_real_file_association(self, tmp_path):
        from spir_dynamic.extraction.xls_converter import convert_xls_to_xlsx
        wb = openpyxl.load_workbook(convert_xls_to_xlsx(REAL_XLS, out_dir=tmp_path), data_only=True)
        result = _extract(wb)
        pumps = [f"{i:04d}" for i in range(1, 34)]
        motors = [f"{i:04d}" for i in range(34, 48)]
        assert list(_by_pos(result, PUMP_A)) == list(_by_pos(result, PUMP_B)) == pumps
        assert list(_by_pos(result, PA_A)) == list(_by_pos(result, PA_B)) == motors
        key = _by_pos(result, PUMP_B)["0021"]
        assert (key["qty_installed_per_unit"], key["total_qty_identical_installed"],
                key["op2y_vendor_pm"], key["op2y_vendor_cm"]) == (8, 16, 5, 5)
        assert _by_pos(result, PA_A)["0047"]["price_unit"] == 29
        assert _by_pos(result, PUMP_A)["0033"]["original_makers_part_no"] == "MTR-EX2515011-005400PH011A/B-S1"


# ── 12-13. Value cleaning and identifiers ───────────────────────────────────

class TestValues:
    def test_cleaning_rules(self):
        line = Line({"pos": "0001", "desc": "  pump shaft\xa0 ", "material": None, "weight": "NA",
                     "unit_price": "-", "c_vendor": 0, "maker": "Abb India Ltd"}, per_unit={0: 1})
        row = _by_pos(_extract(saipem_workbook(lines=[line])), PUMP_A)["0001"]
        assert row["description_of_parts"] == "PUMP SHAFT"     # trimmed (NBSP too), uppercased
        assert row["material"] is None                          # blank stays blank
        assert row["weight"] == "NA"                            # NA stays NA
        assert row["price_unit"] == "-"                         # explicit dash kept
        assert row["cap_vendor"] == 0                           # zero kept
        assert row["original_makers_name"] == "ABB INDIA LTD"

    def test_identifiers(self):
        result = _extract()
        line1, line2 = _by_pos(result, PUMP_A)["0001"], _by_pos(result, PUMP_A)["0002"]
        assert line1["spir_pos_no"] == "0001"
        assert line2["original_makers_part_no"] == "88496189"   # numeric part no → text, no '.0'
        assert line2["rel_doc_pos_no"] == "2120.1"
        assert line1["rel_doc_pos_no"] == "2110"
        assert _by_pos(result, PA_A)["0003"]["original_makers_part_no"] == "M734QK108_CMB"

    def test_zero_padded_number_format_keeps_leading_zeros(self):
        wb = openpyxl.Workbook()
        cell = wb.active["A1"]
        cell.value, cell.number_format = 7, "0000"
        assert identifier_text(cell) == "0007"

    def test_multiline_address_and_phone_text(self):
        result = _extract()
        row = result.rows[0]
        assert row["vendor_address"] == "S.F.NO:136/3,137,MYLERIPALAYAM POST,\nCOIMBATORE641032.TAMILNADU,INDIA."
        assert row["vendor_phone_fax"] == "+91 0422 661 1886"

    def test_numeric_phone_stays_identifier_text(self):
        vendor = {"VENDOR'S NAME\n": "X", "ADDRESS\n": "Y", "PHONE No. / FAX No.\n": 914226611886,
                  "E-MAIL ADDRESS": "a@b.com"}
        assert _extract(saipem_workbook(vendor=vendor)).rows[0]["vendor_phone_fax"] == "914226611886"

    @pytest.mark.parametrize("pm, cm, total", [(1, 1, 2), (5, None, 5), (None, None, None),
                                                (1, "NA", None), (0.5, 1.5, 2), (0, 0, 0)])
    def test_vendor_total(self, pm, cm, total):
        assert vendor_total(pm, cm) == total


# ── 14. Vendor metadata ─────────────────────────────────────────────────────

class TestVendor:
    def test_vendor_fields_on_every_row(self):
        result = _extract()
        assert result.vendor["vendor_name"] == "Flowserve India Controls Pvt Ltd"
        assert {r["vendor_name"] for r in result.rows} == {"FLOWSERVE INDIA CONTROLS PVT LTD"}
        assert {r["vendor_email"] for r in result.rows} == {"MOS@FLOWSERVE.COM"}

    def test_empty_vendor_value_does_not_borrow_neighbour_label(self):
        vendor = {"VENDOR'S NAME\n": None, "ADDRESS\n": None, "PHONE No. / FAX No.\n": None,
                  "E-MAIL ADDRESS": None}
        row = _extract(saipem_workbook(vendor=vendor)).rows[0]
        assert row["vendor_name"] is None and row["vendor_address"] is None
        assert row["vendor_phone_fax"] is None and row["vendor_email"] is None


# ── 15. Dynamic layout ──────────────────────────────────────────────────────

class TestDynamicLayout:
    def test_five_equipment_columns_shift_matrix(self):
        equipment = [Equipment(f"P-{i}00", f"S{i}", f"M{i}", units=1) for i in range(1, 6)]
        lines = [Line({"pos": "0001", "desc": "SEAL", "pm": 1, "cm": 2}, per_unit={0: 1, 4: 3})]
        result = _extract(saipem_workbook(equipment=equipment, lines=lines))
        assert {r["item_tag_no"] for r in result.rows} == {"P-100", "P-500"}
        assert _by_pos(result, "P-500")["0001"]["qty_installed_per_unit"] == 3
        assert _by_pos(result, "P-500")["0001"]["op2y_vendor_total"] == 3

    def test_renamed_band_keeps_its_source_label(self):
        bands = [(" SPARE PARTS IDENTIFICATION\n", SPARE_FIELDS),
                 ("QUANTITY FOR 3 YEARS OPERATION", OPERATING_FIELDS),
                 ("CAPITAL SPARES", CAPITAL_FIELDS), (None, TAIL_FIELDS)]
        result = _extract(saipem_workbook(bands=bands))
        assert result.band_labels["op2y"] == "QUANTITY FOR 3 YEARS OPERATION"
        assert _by_pos(result, PUMP_A)["0002"]["op2y_vendor_pm"] == 5
        assert result.warnings == []

    def test_additional_group_exported_as_extra_columns(self):
        extra = ("COMMISSIONING SPARES", [Field("RECOMMEND. By VENDOR", "comm_vendor"),
                                          Field("PURCHASED QUANTITY\n(approved)", "comm_purch")])
        bands = [(" SPARE PARTS IDENTIFICATION\n", SPARE_FIELDS), ("QUANTITY FOR 2 YEARS OPERATION\n", OPERATING_FIELDS),
                 extra, ("CAPITAL SPARES", CAPITAL_FIELDS), (None, TAIL_FIELDS)]
        lines = two_pump_lines()
        lines[0].values.update(comm_vendor=3, comm_purch=1)
        result = _extract(saipem_workbook(bands=bands, lines=lines))
        extras = [c for c in result.columns if c.id.startswith("extra__")]
        assert [(result.band_labels[c.band], c.field) for c in extras] == [
            ("COMMISSIONING SPARES", "RECOMMEND. By VENDOR"),
            ("COMMISSIONING SPARES", "PURCHASED QUANTITY (approved)")]
        row = _by_pos(result, PUMP_A)["0001"]
        assert [row[c.id] for c in extras] == [3, 1]
        assert row["cap_vendor"] == 0                 # capital group still mapped correctly
        assert sum("not a standard Saipem field" in w for w in result.warnings) == 2
        # the extra group keeps its own heading in the workbook
        ws = openpyxl.load_workbook(io.BytesIO(build_saipem_xlsx(
            result.rows, result.columns, result.band_labels, result.spir_no))).active
        assert "COMMISSIONING SPARES" in [c.value for c in ws[1]]

    def test_missing_field_reported(self):
        spare = [f for f in SPARE_FIELDS if f.key != "vcode"]
        bands = [(" SPARE PARTS IDENTIFICATION\n", spare), ("QUANTITY FOR 2 YEARS OPERATION\n", OPERATING_FIELDS),
                 ("CAPITAL SPARES", CAPITAL_FIELDS), (None, TAIL_FIELDS)]
        result = _extract(saipem_workbook(bands=bands))
        assert any("VENDOR CODE No." in w for w in result.warnings)
        assert all(r.get("vendor_code_no") is None for r in result.rows)

    def test_layout_found_from_labels(self):
        layout = resolve_layout(saipem_workbook()["SPIR Form"])
        assert layout.equipment_cols == [6, 7] and layout.total_col == 8 and layout.pos_col == 9
        assert (layout.data_first, layout.data_last) == (15, 17)
        assert layout.currency == "USD"


# ── 16-17. Excluded generic steps; generic output unchanged ─────────────────

class TestGenericIsolation:
    def test_generic_steps_never_called_for_saipem(self, tmp_path):
        # _run_pipeline patches every generic step to raise; reaching here proves none ran
        path = tmp_path / "form.xlsx"
        saipem_workbook().save(path)
        result, out = _run_pipeline(path, path.name)
        header = {c.value for row in out.active.iter_rows(max_row=3) for c in row if c.value}
        assert not header & {"PLANNING PLANT", "PLANNING PLANT DESCRIPTION", "UNIT PRICE (QAR)",
                             "OLD MATERIAL NUMBER/SPF NUMBER", "ERROR", "S.NO"}
        assert result["dup1_count"] == 0 and result["sap_count"] == 0

    def test_generic_schema_unchanged(self):
        assert OUTPUT_COLS == [
            "SPIR NO", "TAG NO", "EQPT MAKE", "EQPT MODEL", "EQPT SR NO", "EQPT QTY",
            "QUANTITY IDENTICAL PARTS FITTED", "ITEM NUMBER", "POSITION NUMBER",
            "OLD MATERIAL NUMBER/SPF NUMBER", "DESCRIPTION OF PARTS", "NEW DESCRIPTION OF PARTS",
            "DWG NO INCL POSN NO", "MANUFACTURER PART NUMBER", "MATERIAL SPECIFICATION",
            "SUPPLIER/ OCM NAME", "CURRENCY", "UNIT PRICE", "UNIT PRICE (QAR)",
            "DELIVERY TIME IN WEEKS", "MIN MAX STOCK LVLS QTY", "UNIT OF MEASURE", "SAP NUMBER",
            "CLASSIFICATION OF PARTS", "ERROR", "SHEET", "SPIR TYPE", "VENDOR NAME",
            "VENDOR EMAIL1", "VENDOR EMAIL2", "VENDOR CONTACT NO", "VENDOR COUNTRY",
            "PLANNING PLANT", "PLANNING PLANT DESCRIPTION",
        ]

    def test_generic_build_xlsx_unchanged(self):
        from spir_dynamic.services.excel_builder import build_xlsx
        ws = openpyxl.load_workbook(io.BytesIO(build_xlsx([["x"] * len(OUTPUT_COLS)], "S"))).active
        assert [c.value for c in ws[1]] == ["S.NO", *OUTPUT_COLS]
        assert ws.cell(4, 1).value == "0001" and ws.cell(4, 2).value == "X"

    def test_generic_workbook_takes_generic_path(self, tmp_path):
        if not GENERIC_WORKBOOK.exists():
            pytest.skip("generic template not present")
        from spir_dynamic.app.pipeline import run_pipeline
        storage = MagicMock()
        with patch("spir_dynamic.app.pipeline.get_storage", return_value=storage), \
             patch("spir_dynamic.app.pipeline._apply_currency_conversion", return_value=None), \
             patch("spir_dynamic.extraction.saipem.service.run_saipem_pipeline",
                   side_effect=AssertionError("Saipem path must not run")), \
             patch("spir_dynamic.app.pipeline.run_saipem_pipeline",
                   side_effect=AssertionError("Saipem path must not run")):
            result = run_pipeline(GENERIC_WORKBOOK, GENERIC_WORKBOOK.name, NGL)
        assert result["preview_cols"] == OUTPUT_COLS
        assert "SAIPEM" not in result["format"]


# ── NEW DESCRIPTION OF PARTS and the equipment row's description ────────────

class TestDescriptions:
    def test_new_description_concatenates_in_standard_order(self):
        # desc, part no, supplier (makers name), material — as the Standard output
        row = _by_pos(_extract(), PUMP_A)["0001"]
        assert row["new_description_of_parts"] == (
            "PUMP SHAFT,88590797,FLOWSERVE INDIA CONTROLS PVT LTD,A582 TYP 416 SS")
        assert row["description_of_parts"] == "PUMP SHAFT"            # source value untouched

    def test_new_description_skips_blank_and_placeholders(self):
        lines = [Line({"pos": "0001", "desc": "SEAL", "material": "-", "maker": "NA", "part": "P-1"}, per_unit={0: 1}),
                 Line({"pos": "0002", "desc": "RING", "material": None, "maker": "ABB", "part": None}, per_unit={0: 1})]
        by_pos = _by_pos(_extract(saipem_workbook(lines=lines)), PUMP_A)
        assert by_pos["0001"]["new_description_of_parts"] == "SEAL,P-1"
        assert by_pos["0002"]["new_description_of_parts"] == "RING,ABB"

    def test_equipment_row_gets_form_description(self):
        result = _extract()
        for tag in (PUMP_A, PUMP_B, PA_A, PA_B):
            first = next(r for r in result.rows if r["item_tag_no"] == tag)
            assert first.get("spir_pos_no") is None
            assert first["description_of_parts"] == "CENTRIFUGAL PUMPS"
            assert first.get("new_description_of_parts") is None
        # spare rows keep their own part descriptions
        assert _by_pos(result, PUMP_A)["0002"]["description_of_parts"] == "KEY - LINE SHAFT"

    def test_missing_form_description_leaves_blank(self):
        result = _extract(saipem_workbook(equipment_description=None))
        assert next(r for r in result.rows if r["item_tag_no"] == PUMP_A).get("description_of_parts") is None

    def test_new_description_column_follows_description(self):
        ids = [c.id for c in SAIPEM_COLUMNS]
        assert ids[ids.index("description_of_parts") + 1] == "new_description_of_parts"

    @needs_real_xls
    def test_real_file_descriptions(self, tmp_path):
        from spir_dynamic.extraction.xls_converter import convert_xls_to_xlsx
        result = _extract(openpyxl.load_workbook(convert_xls_to_xlsx(REAL_XLS, out_dir=tmp_path), data_only=True))
        assert result.rows[0]["description_of_parts"] == "CENTRIFUGAL PUMPS"
        assert _by_pos(result, PUMP_A)["0033"]["new_description_of_parts"] == (
            "MOTOR BEARING SET(DE / NDE),MTR-EX2515011-005400PH011A/B-S1,ABB INDIA LTD")
