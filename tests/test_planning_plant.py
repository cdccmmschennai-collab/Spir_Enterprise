"""
The user-selected Planning Plant is the single source of truth for the
PLANNING PLANT / PLANNING PLANT DESCRIPTION output columns.

Covers:
    - validation against the controlled 9-plant master (missing / invalid rejected)
    - run_pipeline stamps the selected plant on every row of a real SPIR workbook,
      even when the workbook text names a different plant (DUKHAN)
    - the Celery task refuses to extract without a valid plant (no fallback)
    - batch paths apply the same plant to every file
    - API entry points reject a missing / invalid plant
    - the old keyword-based plant classifier is gone
"""
from __future__ import annotations

import asyncio
import importlib.util
import inspect
import io
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock, patch

import openpyxl
import pytest

from spir_dynamic.extraction.output_schema import CI, OUTPUT_COLS
from spir_dynamic.services.planning_plant import (
    PLANT_MASTER,
    InvalidPlanningPlant,
    PlanningPlant,
    resolve_planning_plant,
)
from tests.test_source_objects import (  # noqa: F401 — `api` is a pytest fixture
    MemoryObjectStorage,
    _slot_updates,
    api,
    worker_ctx,
)

REPO_ROOT = Path(__file__).resolve().parents[1]
# Project title: "EPIC OF DUKHAN PRODUCTION FACILITIES UPGRADE - PHASE 1B" (a 2300 keyword).
DUKHAN_WORKBOOK = REPO_ROOT / "templates" / "inputs" / "VEN-4460-DGEN-5-43-0115-1.xlsm"

NGL = PlanningPlant("2400", "NGL Mesaieed")
DUKHAN = PlanningPlant("2300", "Dukhan Fields")
NGL_FIELDS = {"planning_plant": "2400", "planning_plant_description": "NGL Mesaieed"}


# ── Validation against the controlled master ────────────────────────────────

class TestResolvePlanningPlant:
    def test_master_is_exactly_the_nine_plants(self):
        assert PLANT_MASTER == {
            "2600": "Ras Laffan Offshore Fields",
            "2300": "Dukhan Fields",
            "2500": "Refinery Mesaieed",
            "1200": "Qatar Petroleum - Doha",
            "2800": "Mesaieed Industrial City",
            "2900": "North Field Alpha",
            "2400": "NGL Mesaieed",
            "2700": "Ras Laffan Industrial City",
            "3000": "RL Cooling Water Systems",
        }

    @pytest.mark.parametrize("code", sorted(PLANT_MASTER))
    def test_every_controlled_plant_resolves(self, code):
        assert resolve_planning_plant(code, PLANT_MASTER[code]) == PlanningPlant(code, PLANT_MASTER[code])

    def test_surrounding_whitespace_ignored(self):
        assert resolve_planning_plant(" 2400 ", " NGL Mesaieed ") == NGL

    @pytest.mark.parametrize("code, desc", [
        (None, None), ("", ""), ("   ", "NGL Mesaieed"), (None, "NGL Mesaieed"),   # code missing
        ("2400", None), ("2400", ""),                                               # description missing
    ])
    def test_missing_rejected(self, code, desc):
        with pytest.raises(InvalidPlanningPlant):
            resolve_planning_plant(code, desc)

    @pytest.mark.parametrize("code, desc", [
        ("9999", "NGL Mesaieed"),              # not a controlled code
        ("24", "NGL Mesaieed"),                # partial code
        ("ANNEXURE 1", "NGL Mesaieed"),        # free text
        ("2400", "Dukhan Fields"),             # another plant's description
        ("2400", "NGL MESAIEED"),              # not the controlled spelling
        ("2400", "My own description"),        # arbitrary description
    ])
    def test_invalid_rejected(self, code, desc):
        with pytest.raises(InvalidPlanningPlant):
            resolve_planning_plant(code, desc)


# ── Output columns ───────────────────────────────────────────────────────────

def test_output_columns_are_planning_plant():
    assert OUTPUT_COLS[-2:] == ["PLANNING PLANT", "PLANNING PLANT DESCRIPTION"]
    assert "PLANT" not in OUTPUT_COLS and "PLANT DESCRIPTION" not in OUTPUT_COLS


def test_apply_plant_overwrites_every_row_unconditionally():
    from spir_dynamic.app.pipeline import _apply_plant
    pc, dc = CI["PLANNING PLANT"], CI["PLANNING PLANT DESCRIPTION"]
    rows = [[None] * len(OUTPUT_COLS) for _ in range(3)]
    rows[1][pc], rows[1][dc] = "2300", "Dukhan Fields"         # anything already there loses
    _apply_plant(rows, NGL)
    assert [(r[pc], r[dc]) for r in rows] == [("2400", "NGL Mesaieed")] * 3


# ── run_pipeline on a real SPIR workbook ─────────────────────────────────────

def _run_real_pipeline(plant: PlanningPlant):
    from spir_dynamic.app.pipeline import run_pipeline
    stored: dict = {}
    storage = MagicMock()
    storage.put.side_effect = lambda fid, data, name, **kw: stored.update(xlsx=data)
    with patch("spir_dynamic.app.pipeline.get_storage", return_value=storage), \
         patch("spir_dynamic.app.pipeline._apply_currency_conversion", return_value=None):
        result = run_pipeline(DUKHAN_WORKBOOK, DUKHAN_WORKBOOK.name, plant)
    return result, stored["xlsx"]


def _excel_plant_columns(xlsx: bytes) -> list[tuple]:
    ws = openpyxl.load_workbook(io.BytesIO(xlsx), read_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    pc, dc = header.index("PLANNING PLANT"), header.index("PLANNING PLANT DESCRIPTION")
    # rows 2-3 are the field-name / character-limit metadata rows; data starts at row 4
    return [(r[pc], r[dc]) for r in rows[3:] if any(v is not None for v in r)]


@pytest.mark.skipif(not DUKHAN_WORKBOOK.exists(), reason="sample SPIR templates not available")
class TestPipelineUsesSelectedPlant:
    @pytest.mark.parametrize("plant", [NGL, DUKHAN], ids=["2400", "2300"])
    def test_selected_plant_on_every_row(self, plant):
        result, xlsx = _run_real_pipeline(plant)
        assert result["total_rows"] > 0
        assert result["plant"] == plant.code and result["plant_description"] == plant.description
        pc, dc = CI["PLANNING PLANT"], CI["PLANNING PLANT DESCRIPTION"]
        # the pipeline's existing final step uppercases every string cell
        expected = (plant.code, plant.description.upper())
        assert {(r[pc], r[dc]) for r in result["preview_rows"]} == {expected}
        excel = _excel_plant_columns(xlsx)
        assert len(excel) == result["total_rows"] and set(excel) == {expected}

    def test_conflicting_workbook_keyword_does_not_override(self):
        """The workbook's project title says DUKHAN; the user picked 2400 — 2400 wins."""
        wb = openpyxl.load_workbook(DUKHAN_WORKBOOK, read_only=True, data_only=True)
        try:
            text = " ".join(str(v) for ws in wb.worksheets
                            for row in ws.iter_rows(values_only=True, max_row=150) for v in row if v)
        finally:
            wb.close()
        assert "DUKHAN" in text.upper()

        result, xlsx = _run_real_pipeline(NGL)
        pc = CI["PLANNING PLANT"]
        assert {r[pc] for r in result["preview_rows"]} == {"2400"}
        assert {c for c, _ in _excel_plant_columns(xlsx)} == {"2400"}


def test_pipeline_without_plant_never_reads_the_file(tmp_path):
    from spir_dynamic.app.pipeline import run_pipeline
    missing = tmp_path / "never-opened.xlsx"
    with pytest.raises(InvalidPlanningPlant):
        run_pipeline(missing, missing.name, None)


# ── The old keyword classifier is gone ───────────────────────────────────────

def test_old_keyword_classifier_removed():
    assert importlib.util.find_spec("spir_dynamic.services.plant_classifier") is None
    from spir_dynamic.app import pipeline
    src = inspect.getsource(pipeline)
    assert "plant_classifier" not in src and "classify_plant" not in src


def test_pipeline_plant_comes_only_from_argument():
    """run_pipeline's plant must be a required parameter — no default, no workbook lookup."""
    from spir_dynamic.app.pipeline import run_pipeline
    param = inspect.signature(run_pipeline).parameters["planning_plant"]
    assert param.default is inspect.Parameter.empty


# ── Celery task ──────────────────────────────────────────────────────────────

class TestWorkerPlanningPlant:
    KEY = "job-1_000_big.xlsm"

    def test_plant_reaches_pipeline(self, tmp_path):
        st = MemoryObjectStorage()
        st.put_bytes(self.KEY, b"wb")
        with worker_ctx(tmp_path, st) as (task, _store, pipeline):
            out = task.run(job_id="job-1", file_idx=0, source_key=self.KEY, filename="big.xlsm",
                           planning_plant="2300", planning_plant_description="Dukhan Fields")
        assert out["status"] == "ok"
        assert pipeline.call_args.args[2] == DUKHAN

    @pytest.mark.parametrize("fields", [
        {},                                                                         # missing
        {"planning_plant": "2400"},                                                 # no description
        {"planning_plant": "1234", "planning_plant_description": "NGL Mesaieed"},   # invalid code
    ])
    def test_missing_or_invalid_plant_fails_without_extraction(self, tmp_path, fields):
        st = MemoryObjectStorage()
        st.put_bytes(self.KEY, b"wb")
        with worker_ctx(tmp_path, st) as (task, store, pipeline):
            with patch.object(task, "retry") as retry:
                out = task.run(job_id="job-1", file_idx=0, source_key=self.KEY, filename="big.xlsm", **fields)
        assert out["status"] == "error" and "planning_plant" in out["error"]
        pipeline.assert_not_called()           # no extraction, so no automatic detection either
        retry.assert_not_called()              # not retryable
        assert self.KEY not in st.objects      # source discarded
        assert [u.status for u in _slot_updates(store)] == ["error"]


# ── Batch: one plant for every file ──────────────────────────────────────────

def test_dispatch_applies_plant_to_every_file():
    from spir_dynamic.app.batch_router import _dispatch_celery
    task = MagicMock()
    files = [(f"j_{i:03d}_f{i}.xlsx", f"f{i}.xlsx", 1) for i in range(5)]
    cfg = SimpleNamespace(large_file_threshold_mb=100, giant_file_threshold_mb=500)
    with patch("spir_dynamic.tasks.extraction_tasks.process_file_task", task), \
         patch("spir_dynamic.monitoring.metrics.GIANT_FILES_ROUTED"):
        _dispatch_celery("j", files, cfg, planning_plant=NGL)
    assert task.apply_async.call_count == 5
    assert all(c.kwargs["kwargs"] == NGL_FIELDS for c in task.apply_async.call_args_list)


def test_fallback_batch_applies_plant_to_every_file(tmp_path):
    from spir_dynamic.app.batch_router import _process_batch_from_storage
    from spir_dynamic.services.job_store import JobStore

    storage = MemoryObjectStorage()
    names = [f"f{i}.xlsx" for i in range(5)]
    for i, n in enumerate(names):
        storage.put_bytes(f"job_{i:03d}_{n}", b"x")
    store = JobStore()
    store.create("job", names)
    seen = []

    def fake_pipeline(path, filename, plant):
        seen.append((filename, plant))
        return {"total_rows": 1, "total_tags": 1, "spir_no": "S", "file_id": filename}

    with patch("spir_dynamic.services.source_objects.get_source_storage", return_value=storage), \
         patch("spir_dynamic.app.batch_router.get_job_store", return_value=store), \
         patch("spir_dynamic.app.batch_router.run_pipeline", side_effect=fake_pipeline), \
         patch("spir_dynamic.services.source_objects.scratch_dir", return_value=tmp_path / "scratch"):
        asyncio.run(_process_batch_from_storage(
            "job", [(f"job_{i:03d}_{n}", n, 1) for i, n in enumerate(names)], planning_plant=NGL,
        ))
    assert seen == [(n, NGL) for n in names]
    assert [r.status for r in store.get("job").results] == ["ok"] * 5


# ── API entry points (TestClient needs httpx) ────────────────────────────────

class TestApiRejectsMissingOrInvalidPlant:
    BAD = [
        pytest.param({}, id="missing"),
        pytest.param({"planning_plant": "2400"}, id="no-description"),
        pytest.param({"planning_plant": "7777", "planning_plant_description": "NGL Mesaieed"}, id="bad-code"),
        pytest.param({"planning_plant": "2400", "planning_plant_description": "Dukhan Fields"}, id="mismatch"),
    ]

    @pytest.mark.parametrize("fields", BAD)
    def test_batch_register(self, api, fields):
        with patch("spir_dynamic.app.batch_router._persist_job_to_db", new=AsyncMock()):
            res = api.client.post("/api/batch/register", json={"filenames": ["a.xlsx"], **fields})
        assert res.status_code == 422

    @pytest.mark.parametrize("fields", BAD)
    def test_batch_extract(self, api, fields):
        with patch("spir_dynamic.app.batch_router._dispatch_celery") as dispatch:
            res = api.client.post("/api/batch/extract", data=fields,
                                  files=[("files", ("a.xlsx", b"1", "application/octet-stream"))])
        assert res.status_code == 422
        dispatch.assert_not_called()
        assert api.storage.objects == {}

    @pytest.mark.parametrize("fields", BAD)
    def test_batch_upload_marks_slot_error(self, api, fields):
        from spir_dynamic.services.job_store import JobStore
        store = JobStore()
        store.create("job", ["a.xlsx"], user_id="user-1")
        with patch("spir_dynamic.app.batch_router.get_job_store", return_value=store), \
             patch("spir_dynamic.app.batch_router._dispatch_celery") as dispatch:
            res = api.client.post("/api/batch/job/upload", data={"file_idx": "0", **fields},
                                  files={"file": ("a.xlsx", b"1", "application/octet-stream")})
        assert res.status_code == 422
        dispatch.assert_not_called()
        assert store.get("job").results[0].status == "error"          # poller never stalls
        assert api.storage.objects == {}

    @pytest.mark.parametrize("fields", BAD)
    def test_direct_upload_initiate(self, api, fields):
        res = api.client.post("/api/uploads/initiate", json={"filename": "a.xlsm", "size": 10, **fields})
        assert res.status_code == 422

    def test_batch_of_five_files_all_get_the_selected_plant(self, api):
        from spir_dynamic.services.job_store import JobStore
        store = JobStore()
        names = [f"f{i}.xlsx" for i in range(5)]
        task = MagicMock()
        with patch("spir_dynamic.app.batch_router.get_settings", return_value=api.cfg), \
             patch("spir_dynamic.app.batch_router.get_job_store", return_value=store), \
             patch("spir_dynamic.app.batch_router._persist_job_to_db", new=AsyncMock()), \
             patch("spir_dynamic.tasks.extraction_tasks.process_file_task", task), \
             patch("spir_dynamic.monitoring.metrics.GIANT_FILES_ROUTED"):
            job_id = api.client.post("/api/batch/register", json={"filenames": names, **NGL_FIELDS}).json()["job_id"]
            for i, n in enumerate(names):
                r = api.client.post(f"/api/batch/{job_id}/upload", data={"file_idx": str(i), **NGL_FIELDS},
                                    files={"file": (n, b"x", "application/octet-stream")})
                assert r.status_code == 200
        assert task.apply_async.call_count == 5
        assert all(c.kwargs["kwargs"] == NGL_FIELDS for c in task.apply_async.call_args_list)
