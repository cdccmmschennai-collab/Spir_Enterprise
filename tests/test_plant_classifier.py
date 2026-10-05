"""
Tests for services/plant_classifier.py — workbook-level Plant classification.

Workbooks are built in memory with openpyxl. Label positions are varied on
purpose: the classifier must locate PROJECT/PLANT by structure, never by a
fixed coordinate such as I32/J32.
"""
from __future__ import annotations

import openpyxl
import pytest

from spir_dynamic.services.plant_classifier import (
    PLANT_MASTER,
    STATUS_CONFLICT,
    STATUS_RESOLVED,
    STATUS_UNRESOLVED,
    classify_plant,
)


def _wb(*sheets: dict[str, object], hidden: tuple[int, ...] = ()) -> openpyxl.Workbook:
    """Build a workbook; each dict maps cell coordinate -> value for one sheet."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for i, cells in enumerate(sheets):
        ws = wb.create_sheet(f"MAIN SHEET {i + 1}")
        for coord, value in cells.items():
            ws[coord] = value
        if i in hidden:
            ws.sheet_state = "hidden"
    return wb


def _main(project: object, label_at: str = "I32", value_at: str = "J32", **extra) -> dict:
    cells = {label_at: "PROJECT/PLANT  :", value_at: project,
             "K32": "MANUFACTURERS/SUPPLIERS FOCAL POINT INCLUDING \nE-MAIL/ TEL/FAX NO:"}
    cells.update(extra)
    return cells


def _resolved(wb, code: str):
    res = classify_plant(wb)
    assert res.status == STATUS_RESOLVED, res.reason
    assert res.plant == code
    assert res.plant_description == PLANT_MASTER[code]
    return res


# ── Real-sample project titles (verbatim from the investigation) ─────────────

@pytest.mark.parametrize("title, code", [
    ("EPIC FOR DUKHAN PRODUCTION FACILITIES UPGRADE (DPFU) - PHASE 1B", "2300"),
    ("DUKHAN FIELDS, FEED FOR DUKHAN PRODUCTION FACILITIES UPGRADE –PHASE 1B", "2300"),
    ("DPFU Phase1B Project", "2300"),
    ("DUKHAN - JALEHA MAIN", "2300"),
    ("EPIC FOR DUKHAN PRODUCTION FACILITIES \nUPGRADE (DPFU)- PHASE 1B", "2300"),
    ("EPIC OF NEW EFFLUENT WATER TREATMENT PLANT FOR NGL AT MESAIEED", "2400"),
    ("EPIC FOR NEW EWTP FOR NGL AT MESAIEED", "2400"),
    ("EPIC FOR  COMMON COOLING SEAWATER SYSTEM  PHASE-3 PROJECT (RLIC)", "3000"),
    ("EPIC FOR BERTH 6 FIRE PROTECTION WORKS AT MIC PORT", "2800"),
])
def test_sample_titles(title, code):
    _resolved(_wb(_main(title)), code)


def test_label_with_inline_value_anywhere():
    wb = _wb({"G31": "PROJECT/PLANT : EPIC FOR  COMMON COOLING SEAWATER SYSTEM  PHASE-3 PROJECT (RLIC)",
              "G32": "CONTRACT No.    : GC19103600"})
    res = _resolved(wb, "3000")
    assert res.evidence[0].cell == "G31"


def test_spaced_label_and_moved_position():
    wb = _wb({"E33": "PROJECT/ PLANT    :  COMMON COOLING SEAWATER SYSTEM PHASE-3 PROJECT    (RLIC)"})
    _resolved(wb, "3000")
    _resolved(_wb(_main("DPFU 1B", label_at="K28", value_at="L28")), "2300")


def test_multiline_label_cell_stops_at_next_label():
    wb = _wb({"H29": "PROJECT/PLANT : EPIC FOR WELL FLOWLINES IN DUKHAN\nCONTRACT No. : GC191008B0\n"
                     "COMPANY : QATAR PETROLEUM"})
    res = _resolved(wb, "2300")
    assert "QATAR PETROLEUM" not in res.evidence[0].text


# ── NGL numbered units ───────────────────────────────────────────────────────

@pytest.mark.parametrize("title", [
    "NGL1 MESAIEED", "NGL 1 MESAIEED", "NGL-1 MESAIEED",
    "NGL2 MESAIEED", "NGL 2 MESAIEED", "NGL-2 MESAIEED", "NGL-1/2 AT MESAIEED",
])
def test_ngl_numbered_units_with_mesaieed(title):
    _resolved(_wb(_main(title)), "2400")


def test_ngl_equipment_location_alone_does_not_classify():
    wb = _wb(_main(None, D4="DCS SYSTEM PANEL NGL-1/2 LER-41 (INFI90)"))
    assert classify_plant(wb).status == STATUS_UNRESOLVED
    wb = _wb(_main("DCS UPGRADE FOR NGL-3"))  # NGL without Mesaieed context
    assert classify_plant(wb).status == STATUS_UNRESOLVED


# ── Approved decisions ───────────────────────────────────────────────────────

def test_rlic_alone_is_unresolved():
    assert classify_plant(_wb(_main("NEW SUBSTATION AT RLIC"))).status == STATUS_UNRESOLVED


def test_rlcsf_in_project_title():
    _resolved(_wb(_main("RLCSF3 PACKAGE")), "3000")


def test_mic_outside_project_context_is_ignored():
    wb = _wb(_main("SOME UNKNOWN PROJECT", I33="COMPANY :", J33="MIC TRADING",
                   I28="MIC SENSOR ASSEMBLY"))
    assert classify_plant(wb).status == STATUS_UNRESOLVED


def test_empty_project_ignores_annexure_and_neighbour_label():
    wb = _wb(_main(None),
             {"A1": "SITE", "B1": "Location", "A2": "Fahahil Main", "B2": "Jaleha Main"})
    res = classify_plant(wb)
    assert res.status == STATUS_UNRESOLVED
    assert res.plant is None and res.plant_description is None
    # The MANUFACTURERS label is not taken as a value; the empty label is audited.
    assert [(e.cell, e.text) for e in res.evidence] == [("I32", "")]
    assert res.reason == "project-level label(s) found but empty"


# ── Hard exclusions ──────────────────────────────────────────────────────────

def test_codes_never_substring_matched():
    wb = _wb(_main("UNKNOWN WORKS", W36="REQUISITION No:", Y36="PO2400550",
                   C32="4500111818", A1=2400, B5="2300"))
    assert classify_plant(wb).status == STATUS_UNRESOLVED


@pytest.mark.parametrize("title", [
    "WORKS AT MESAIEED", "WORKS AT RAS LAFFAN", "DOHA OFFICE WORKS", "QATAR PETROLEUM", "QATARENERGY",
])
def test_generic_identifiers_do_not_classify(title):
    assert classify_plant(_wb(_main(title))).status == STATUS_UNRESOLVED


def test_company_and_vendor_blocks_are_ignored():
    wb = _wb(_main(None, I33="COMPANY           :", J33="QATAR PETROLEUM",
                   K33="DOHA ENGINEERING SERVICES CO. W.L.L. P.O. Box 14172, Doha, Qatar"))
    assert classify_plant(wb).status == STATUS_UNRESOLVED


def test_hidden_sheets_are_ignored():
    wb = _wb(_main("SOME UNKNOWN PROJECT"), _main("EPIC FOR DUKHAN WORKS"), hidden=(1,))
    assert classify_plant(wb).status == STATUS_UNRESOLVED


# ── Explicit codes and exact descriptions ────────────────────────────────────

@pytest.mark.parametrize("value", [2900, "2900", 2900.0])
def test_explicit_code_label(value):
    res = _resolved(_wb({"B4": "PLANT CODE :", "C4": value}), "2900")
    assert res.rule == "explicit_code"


def test_explicit_code_label_rejects_non_exact_values():
    assert classify_plant(_wb({"B4": "MPP :", "C4": "29001"})).status == STATUS_UNRESOLVED
    assert classify_plant(_wb({"B4": "PLANT CODE :", "C4": "1234"})).status == STATUS_UNRESOLVED


def test_mpp_reference_table_does_not_classify():
    cells = {"A1": "MPP", "B1": "MPP Description"}
    for i, (code, desc) in enumerate(PLANT_MASTER.items(), start=2):
        cells[f"A{i}"], cells[f"B{i}"] = int(code), desc
    assert classify_plant(_wb(cells)).status == STATUS_UNRESOLVED


@pytest.mark.parametrize("code", sorted(PLANT_MASTER))
def test_exact_description_for_every_plant(code):
    res = _resolved(_wb(_main(f"EPIC WORKS - {PLANT_MASTER[code]}")), code)
    assert res.rule == f"exact_description:{code}"


# ── Consistency checking ─────────────────────────────────────────────────────

def test_conflicting_sheets_are_conflict():
    wb = _wb(_main("EPIC FOR DUKHAN WORKS"), _main("EPIC FOR NGL AT MESAIEED"))
    res = classify_plant(wb)
    assert res.status == STATUS_CONFLICT
    assert res.plant is None and res.plant_description is None


def test_conflict_within_one_title():
    res = classify_plant(_wb(_main("NGL AT MESAIEED INDUSTRIAL CITY")))
    assert res.status == STATUS_CONFLICT


def test_consistent_sheets_and_empty_sheet_resolve():
    wb = _wb(_main("DPFU 1B"), _main("EPIC FOR DUKHAN PRODUCTION FACILITIES UPGRADE"), _main(None))
    res = _resolved(wb, "2300")
    assert [e.plants for e in res.evidence] == [["2300"], ["2300"], []]
