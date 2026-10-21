"""
Synthetic Saipem SPIR form workbooks for tests (no customer data).

build_form() reproduces the structure of the Saipem 'SPIR Form' sheet:
merged equipment label column + code column, one column per equipment group,
the TOTAL QUANTITY OF IDENTICAL PARTS INSTALLED column, a banded 3-row matrix
header built from merged ranges, the merged per-unit label spanning the spare
rows, a SUB-TOTAL footer and the labelled header block (SPIR No., REV, vendor).

Everything is relative to the number of equipment columns, so moving the
matrix sideways (as the 5-group Example sheet does) is just equipment=[…5…].
"""
from __future__ import annotations

from dataclasses import dataclass, field
from typing import Any, Optional

import openpyxl


@dataclass
class Equipment:
    tag: Any
    serial: Any = None
    model: Any = None
    free_field: Any = "-"
    sub_item: Any = "-"
    units: Any = None


@dataclass
class Field:
    label: str
    key: str                                    # key into Line.values
    subs: Optional[list[tuple[str, str]]] = None  # (sub label, key)
    qualifier: Optional[str] = None             # row-13 label spanning the subs (PRICE → CURRENCY)


@dataclass
class Line:
    values: dict[str, Any]
    per_unit: dict[int, Any] = field(default_factory=dict)   # equipment index → qty
    total: Any = None


SPARE_FIELDS = [
    Field("SPIR POS. No. \n(from 0001)", "pos"),
    Field("DESCRIPTION OF PARTS", "desc"),
    Field("MATERIAL ", "material"),
    Field("ORIGINAL MAKERS NAME", "maker"),
    Field("ORIGINAL MAKERS PART No.", "part"),
    Field("RELEVANT DOCUMENT No.\n", "doc",
          subs=[("DRAWING No.", "dwg"), ("PAGE", "page"), ("POS. No.\n", "dpos")]),
    Field("VENDOR\nCODE No", "vcode"),
    Field("Business unit codification (Miap code)", "miap"),
]
OPERATING_FIELDS = [
    Field("UNIT of MEAS.", "uom"),
    Field("RECOMMEND. By VENDOR", "vendor",
          subs=[("Quantity for PM", "pm"), ("Quantity for CM", "cm")]),
    Field("ENGINEER. RECOMMEND. ", "eng"),
    Field("OPERATION RECOMMEND. ", "op"),
    Field("PURCHASED QUANTITY\n(approved)", "purch"),
]
CAPITAL_FIELDS = [
    Field("RECOMMEND. By VENDOR", "c_vendor"),
    Field("ENGINEER. RECOMMEND. ", "c_eng"),
    Field("OPERATION RECOMMEND. ", "c_op"),
    Field("PURCHASED QUANTITY\n(approved)", "c_purch"),
]
TAIL_FIELDS = [
    Field("WEIGHT", "weight"),
    Field("CHARACTERISTIC DIMENSION", "dim"),
    Field("OBSOLESCENCE", "obs"),
    Field("G/S", "gs"),
    Field("SHELF LIFE", "shelf"),
    Field("PRICE ", "price", subs=[("UNIT", "unit_price"), ("AMOUNT", "amount")],
          qualifier="CURRENCY\n(USD )"),
    Field("DELIVERY TIME\n(IN WEEKS)", "delivery"),
]


def default_bands() -> list[tuple[Optional[str], list[Field]]]:
    return [
        (" SPARE PARTS IDENTIFICATION\n", SPARE_FIELDS),
        ("QUANTITY FOR 2 YEARS OPERATION\n", OPERATING_FIELDS),
        ("CAPITAL SPARES", CAPITAL_FIELDS),
        (None, TAIL_FIELDS),
    ]


def two_pump_equipment() -> list[Equipment]:
    return [
        Equipment("005400PH011A/B", "1825076/1825077", "08ELM", units=2),
        Equipment("004300PA001A/B", "EX13251310 / EX13251309", "1K40-25-125A-OP", units=2),
    ]


def two_pump_lines() -> list[Line]:
    """Lines 0001-0002 fit only group 0, line 0003 only group 1 (as in the real form)."""
    common = {"material": "A582 TYP 416 SS", "maker": "Flowserve India Controls Pvt Ltd",
              "dwg": "102700DVDPFW607", "page": 3, "uom": "EA", "c_vendor": 0,
              "weight": "NA", "dim": "NA", "obs": 10, "gs": "S", "shelf": 5, "delivery": 24}
    return [
        Line({**common, "pos": "0001", "desc": "PUMP SHAFT ", "part": "88590797", "dpos": 2110,
              "vcode": "88590797", "pm": 1, "cm": 1, "unit_price": "-"},
             per_unit={0: 1}, total=2),
        Line({**common, "pos": "0002", "desc": "KEY - LINE SHAFT", "part": 88496189, "dpos": 2120.1,
              "vcode": 88496189, "pm": 5, "cm": 5, "unit_price": "-"},
             per_unit={0: 8}, total=16),
        Line({**common, "pos": "0003", "desc": "SHAFT WITH COUPLING KEY",
              "part": "\xa0 M734QK108_CMB", "dpos": "4590.1/4590.2", "vcode": "\xa0 M734QK108_CMB",
              "pm": 1, "cm": 1, "unit_price": 2408, "amount": 0, "delivery": 16},
             per_unit={1: 1}, total=2),
    ]


def build_form(
    ws,
    equipment: list[Equipment],
    lines: list[Line],
    bands: Optional[list[tuple[Optional[str], list[Field]]]] = None,
    spir_no: Any = "102700DKEZFW657",
    rev: Any = 4,
    vendor: Optional[dict[str, Any]] = None,
    revision_block: Any = None,
    equipment_description: Any = "CENTRIFUGAL PUMPS",
) -> None:
    bands = bands if bands is not None else default_bands()
    vendor = vendor if vendor is not None else {
        "VENDOR'S NAME\n": "Flowserve India Controls Pvt Ltd",
        "ADDRESS\n": "S.F.NO:136/3,137,MYLERIPALAYAM POST,\nCOIMBATORE641032.TAMILNADU,INDIA.",
        "PHONE No. / FAX No.\n": "+91 0422 661 1886",
        "E-MAIL ADDRESS": "MoS@flowserve.com",
    }

    def put(r1, c1, value, r2=None, c2=None):
        ws.cell(r1, c1, value=value)
        if (r2 or r1, c2 or c1) != (r1, c1):
            ws.merge_cells(start_row=r1, start_column=c1, end_row=r2 or r1, end_column=c2 or c1)

    # Equipment label + code column (C:E)
    put(2, 3, "EQUIPMENT IDENTIFICATION\n", 13, 3)
    for (label, code, r1, r2) in [
        ("ITEM TAG No. (MAIN ITEM)\n", "A", 2, 4),
        ("Free Field (Second TAG Codification or Item Criticality)", "B", 5, 6),
        ("SERIAL No.\n", "C", 7, 8),
        ("MODEL OR TYPE\n", "D", 9, 10),
        ("EQUIPMENT SUB ITEM No.\n", "E", 11, 13),
    ]:
        put(r1, 4, label, r2, 4)
        put(r1, 5, code, r2, 5)
    put(14, 3, "No. OF UNITS ", 14, 4)
    ws.cell(14, 5, value="F")

    # Equipment columns
    eq_cols = []
    for i, eq in enumerate(equipment):
        c = 6 + i
        eq_cols.append(c)
        for value, r1, r2 in [(eq.tag, 2, 4), (eq.free_field, 5, 6), (eq.serial, 7, 8),
                              (eq.model, 9, 10), (eq.sub_item, 11, 13)]:
            put(r1, c, value, r2, c)
        ws.cell(14, c, value=eq.units)

    total_col = 6 + len(equipment)
    put(2, total_col, " TOTAL QUANTITY OF IDENTICAL PARTS INSTALLED\n", 13, total_col)
    pos_col = total_col + 1

    # Matrix header: band row 10, header rows 11-13
    col_of: dict[str, int] = {}
    c = pos_col
    for band_label, fields in bands:
        band_start = c
        for f in fields:
            if f.subs:
                n = len(f.subs)
                if f.qualifier:
                    put(11, c, f.label, 11, c + n - 1)
                    for j, (sub, key) in enumerate(f.subs):
                        put(12, c + j, sub)
                        col_of[key] = c + j
                    put(13, c, f.qualifier, 13, c + n - 1)
                else:
                    put(11, c, f.label, 12, c + n - 1)
                    for j, (sub, key) in enumerate(f.subs):
                        put(13, c + j, sub)
                        col_of[key] = c + j
                c += n
            else:
                put(11, c, f.label, 13, c)
                col_of[f.key] = c
                c += 1
        if band_label:
            put(10, band_start, band_label, 10, c - 1)
    last_col = c - 1

    # Spare rows: per-unit label merged over exactly the spare rows
    first = 15
    last = first + max(len(lines), 1) - 1
    put(first, 3, "QUANTITY OF PARTS INSTALLED PER ONE UNIT OF EQUIPMENT SHOWN IN  COLUMN A-E\n", last, 4)
    put(first, 5, "G", last, 5)
    for k, line in enumerate(lines):
        r = first + k
        for key, value in line.values.items():
            if key in col_of:
                ws.cell(r, col_of[key], value=value)
        for idx, qty in line.per_unit.items():
            ws.cell(r, eq_cols[idx], value=qty)
        ws.cell(r, total_col, value=line.total)
    if "unit_price" in col_of:
        ws.cell(last + 1, col_of["unit_price"], value="SUB-TOTAL")

    # Header block: SPIR No. / REV top-right, vendor block in the middle
    put(2, pos_col, 'SPARE PARTS INTERCHANGEABILITY RECORD "SPIR"', 2, pos_col + 10)
    put(2, last_col - 7, "SPIR No.")
    put(2, last_col - 5, spir_no, 2, last_col - 2)
    put(2, last_col - 1, "REV")
    ws.cell(2, last_col, value=rev)
    put(4, pos_col, "OWNER\n", 4, pos_col + 1)
    put(4, pos_col + 2, "LIVERPOOL BAY CCS LTD", 4, pos_col + 5)
    for row, (label, value) in zip((4, 5, 8, 9), vendor.items()):
        put(row, pos_col + 6, label, row, pos_col + 8)
        put(row, pos_col + 9, value, row, pos_col + 15)
    put(4, pos_col + 18, "VENDOR'S ISSUE DATE\n", 4, pos_col + 20)
    if revision_block is not None:
        put(last + 3, pos_col + 10, "REVISION No.", last + 3, pos_col + 20)
        put(last + 3, pos_col + 21, revision_block, last + 3, pos_col + 24)
    # Revision block DESCRIPTION (equipment description), as in the real form
    put(last + 4, pos_col + 10, "DESCRIPTION ", last + 4, pos_col + 20)
    put(last + 4, pos_col + 21, equipment_description, last + 4, pos_col + 24)


def build_instruction(ws) -> None:
    """The Instruction sheet: the same labels, listed down one column."""
    rows = [
        ("EQUIPMENT IDENTIFICATION", "Item Tag No. (Main Item)", "Column A"),
        (None, "Free Field", "Column B"),
        (None, "Serial No.", "Column C"),
        (None, "Model or Type", "Column D"),
        (None, "Equipment Sub Item No.", "Column E"),
        (None, "No. of Units", "Column F"),
        (None, "Quantity of parts installed per one unit of equipment", "Column G"),
        (None, "Total quantity of identical parts installed", "Column I"),
        (" SPARE PARTS IDENTIFICATION", "SPIR pos. No.", "Column J"),
        (None, "Description of parts", "Column K"),
    ]
    ws.cell(4, 1, value="Section")
    ws.cell(4, 2, value="Field Label")
    for i, (a, b, c) in enumerate(rows, start=5):
        ws.cell(i, 1, value=a)
        ws.cell(i, 2, value=b)
        ws.cell(i, 3, value=c)


def saipem_workbook(
    equipment: Optional[list[Equipment]] = None,
    lines: Optional[list[Line]] = None,
    with_example: bool = True,
    with_instruction: bool = True,
    **form_kwargs,
) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SPIR Form"
    build_form(ws, equipment if equipment is not None else two_pump_equipment(),
               lines if lines is not None else two_pump_lines(), **form_kwargs)
    if with_example:
        example = wb.create_sheet("SPIR Form Example")
        build_form(
            example,
            [Equipment("220PA001A Pump", "XX1 ", "AAA", units=1), Equipment("220PA001B Pump", " XX2", "AAA", units=1),
             Equipment("230PA002A Pump", "YY1", "BBB", units=1), Equipment("230PA002A Pump", "YY2", "BBB", units=1),
             Equipment("360VA001Z Ball Valve", "ZZAA", "CCC", units=5)],
            [Line({"pos": "0001", "desc": "Bearing", "part": "RR004"}, per_unit={0: 2, 1: 2}, total=4)],
            spir_no="V-", rev=2,
        )
    if with_instruction:
        build_instruction(wb.create_sheet("Instruction"))
    return wb
